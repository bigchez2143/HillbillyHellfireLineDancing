"""Local user-owned move records, immutable snapshots and reviewed imports.

The common reference catalog and core generator are never modified here. SQLite
transactions protect batches and version checks; media originals are immutable.
No URL is fetched and no AI service is called. AUTHOR_REVIEWED is a user's own
attestation, not instructor certification or generator eligibility.
"""
from contextlib import contextmanager
from copy import deepcopy
from datetime import datetime, timezone, timedelta
import csv
import hashlib
import io
import json
import os
from pathlib import Path, PurePosixPath
import re
import sqlite3
import uuid
from urllib.parse import urlsplit, parse_qsl
import zipfile

from .choreography import compile_choreography, exact_count

LIBRARY_DIR = None  # tests may override; otherwise environment is read per call
MAX_IMPORT_BYTES = 8 * 1024 * 1024
MAX_IMAGE_BYTES = 20 * 1024 * 1024
MAX_VIDEO_BYTES = 100 * 1024 * 1024
MAX_ARCHIVE_BYTES = 32 * 1024 * 1024
MAX_ROWS = 1000
FIELDS = {'name', 'aliases', 'explanation', 'duration_counts', 'difficulty',
          'events', 'origins', 'links'}
DIFFICULTIES = {'Unspecified', 'Absolute Beginner', 'Beginner', 'Improver',
                'Intermediate', 'Advanced'}
SECRET_KEYS = {'api_key', 'apikey', 'access_token', 'refresh_token', 'password',
               'authorization', 'client_secret'}


class Conflict(ValueError):
    pass


def _now():
    return datetime.now(timezone.utc).isoformat()


def _json(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(',', ':'), allow_nan=False)


def _id(prefix):
    return prefix + '-' + uuid.uuid4().hex


def _check_id(value, prefix):
    if not isinstance(value, str) or not re.fullmatch(prefix + r'-[0-9a-f]{32}', value):
        raise ValueError('Invalid ' + prefix + ' identifier.')
    return value


def directory():
    if LIBRARY_DIR is not None:
        return Path(LIBRARY_DIR)
    configured = os.environ.get('LINE_DANCE_LIBRARY_DIR')
    if configured:
        return Path(configured).expanduser()
    data_root = os.environ.get('LINE_DANCE_DATA_DIR')
    if data_root:
        return Path(data_root).expanduser() / 'library'
    base = Path(os.environ.get('LOCALAPPDATA', str(Path.home() / '.local' / 'share')))
    return base / 'LineDanceCreator' / 'library'


@contextmanager
def _db(write=False):
    root = directory()
    root.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(root / 'library.sqlite3', timeout=15)
    conn.row_factory = sqlite3.Row
    try:
        conn.executescript('''
          CREATE TABLE IF NOT EXISTS records(id TEXT PRIMARY KEY,version INTEGER NOT NULL,payload TEXT NOT NULL);
          CREATE TABLE IF NOT EXISTS versions(record_id TEXT,version INTEGER,payload TEXT NOT NULL,PRIMARY KEY(record_id,version));
          CREATE TABLE IF NOT EXISTS previews(id TEXT PRIMARY KEY,digest TEXT,payload TEXT,created TEXT,result TEXT);
          CREATE TABLE IF NOT EXISTS batches(id TEXT PRIMARY KEY,changes TEXT,state TEXT);
          CREATE TABLE IF NOT EXISTS media(id TEXT PRIMARY KEY,record_id TEXT,filename TEXT,payload TEXT);
        ''')
        if write:
            conn.execute('BEGIN IMMEDIATE')
        yield conn
        conn.commit()
    except BaseException:
        conn.rollback()
        raise
    finally:
        conn.close()


def _no_secrets(value):
    if isinstance(value, dict):
        for key, child in value.items():
            if str(key).casefold().replace('-', '_') in SECRET_KEYS:
                raise ValueError('Remove credentials from library data before importing.')
            _no_secrets(child)
    elif isinstance(value, list):
        for child in value:
            _no_secrets(child)


def _text(value, label, limit=10000, required=False):
    if not isinstance(value, str):
        raise ValueError(label + ' must be text.')
    value = value.strip()
    if len(value) > limit or '\x00' in value or (required and not value):
        raise ValueError(label + ' is missing or too long.')
    return value


def _url(value):
    value = _text(value, 'URL', 3000, True)
    try:
        parsed = urlsplit(value)
        if parsed.scheme not in {'http', 'https'} or not parsed.hostname or parsed.username or parsed.password:
            raise ValueError('Use an HTTP or HTTPS link without embedded credentials.')
        if any(k.casefold().replace('-', '_') in SECRET_KEYS for k, _ in parse_qsl(parsed.query)):
            raise ValueError('Remove credentials from the URL.')
    except (ValueError, UnicodeError) as exc:
        raise ValueError('Use an HTTP or HTTPS link without embedded credentials.') from exc
    return value


def _count(value, label, optional=False, signed=False):
    if optional and value in (None, ''):
        return None
    try:
        result = exact_count(value)
    except (ValueError, TypeError, ZeroDivisionError, OverflowError) as exc:
        raise ValueError(label + ' must be a finite number or fraction.') from exc
    if abs(result) > 100000 or (not signed and result < 0):
        raise ValueError(label + ' is out of range.')
    return str(result)


def _fields(data):
    if not isinstance(data, dict):
        raise ValueError('Move fields must be an object.')
    _no_secrets(data)
    unknown = set(data) - FIELDS
    if unknown:
        raise ValueError('Unknown move fields: ' + ', '.join(sorted(unknown)))
    aliases = data.get('aliases', [])
    if isinstance(aliases, str):
        aliases = [a.strip() for a in aliases.split(';') if a.strip()]
    if not isinstance(aliases, list) or len(aliases) > 50:
        raise ValueError('Use at most 50 aliases.')
    aliases = list(dict.fromkeys(_text(a, 'Alias', 200, True) for a in aliases))
    level = data.get('difficulty') or 'Unspecified'
    if level not in DIFFICULTIES:
        raise ValueError('Choose a supported difficulty.')
    events = data.get('events') or []
    if not isinstance(events, list) or len(events) > 512:
        raise ValueError('Use an event list with at most 512 entries.')
    allowed = {'id', 'offset_counts', 'duration_counts', 'text', 'support_before',
               'support_after', 'rotation_deg', 'facing_before_deg', 'facing_after_deg'}
    normalized = []
    for event in events:
        if not isinstance(event, dict) or set(event) - allowed:
            raise ValueError('An event contains unsupported fields.')
        event = deepcopy(event)
        for key in ('id', 'text'):
            if key in event:
                event[key] = _text(event[key], key, 10000 if key == 'text' else 200)
        for key in ('offset_counts', 'duration_counts', 'rotation_deg', 'facing_before_deg', 'facing_after_deg'):
            if key in event:
                event[key] = _count(event[key], key, signed=key.endswith('_deg'))
        for key in ('support_before', 'support_after'):
            supports = {'L', 'R', 'both', 'neither', 'unknown'} | ({'any'} if key == 'support_before' else {'same'})
            if key in event and event[key] not in supports:
                raise ValueError('Invalid support state.')
        normalized.append(event)
    origins = data.get('origins', [])
    links = data.get('links', [])
    if not isinstance(origins, list) or len(origins) > 50 or not isinstance(links, list) or len(links) > 50:
        raise ValueError('Use at most 50 origins and 50 links.')
    clean_origins, clean_links = [], []
    for origin in origins:
        if not isinstance(origin, dict) or set(origin) - {'kind', 'label', 'url', 'locator', 'rights_note'}:
            raise ValueError('Invalid origin fields.')
        clean_origins.append({k: _url(v) if k == 'url' else _text(v, k, 3000) for k, v in origin.items()})
    for link in links:
        if isinstance(link, str):
            link = {'url': link}
        if not isinstance(link, dict) or set(link) - {'url', 'label'} or 'url' not in link:
            raise ValueError('Links require a URL and optional label.')
        clean_links.append({'url': _url(link['url']), 'label': _text(link.get('label', ''), 'Link label', 300)})
    return {'name': _text(data.get('name', ''), 'Name', 200, True), 'aliases': aliases,
            'explanation': _text(data.get('explanation', ''), 'Explanation'),
            'duration_counts': _count(data.get('duration_counts'), 'Duration', optional=True),
            'difficulty': level, 'events': normalized, 'origins': clean_origins, 'links': clean_links}


def _get(conn, record_id, version=None):
    _check_id(record_id, 'move')
    if version is None:
        row = conn.execute('SELECT payload FROM records WHERE id=?', (record_id,)).fetchone()
    else:
        if type(version) is not int or version < 1:
            raise ValueError('Version must be a positive integer.')
        row = conn.execute('SELECT payload FROM versions WHERE record_id=? AND version=?', (record_id, version)).fetchone()
    if row is None:
        raise FileNotFoundError('Move version not found.')
    return json.loads(row['payload'])


def _expect(record, version):
    if type(version) is not int or version != record['version']:
        raise Conflict('This move changed. Reload before saving; current version is ' + str(record['version']) + '.')


def _put(conn, record):
    payload = _json(record)
    conn.execute('INSERT INTO versions VALUES(?,?,?)', (record['id'], record['version'], payload))
    conn.execute('INSERT INTO records VALUES(?,?,?) ON CONFLICT(id) DO UPDATE SET version=excluded.version,payload=excluded.payload',
                 (record['id'], record['version'], payload))
    return deepcopy(record)


def _new(fields):
    now = _now()
    return {**deepcopy(fields), 'id': _id('move'), 'version': 1, 'created_at': now,
            'updated_at': now, 'attachments': [], 'review': {'status': 'UNVERIFIED'},
            'deleted': False, 'generator_eligible': False}


def _next(record):
    record = deepcopy(record)
    record['version'] += 1
    record['updated_at'] = _now()
    return record


def reference_catalog():
    from .steps import MOVES
    raw = json.loads((Path(__file__).resolve().parents[2] / 'data' / 'step-database.json').read_text(encoding='utf-8-sig'))
    reference = [{**row, 'id': 'reference-' + str(index), 'read_only': True} for index, row in enumerate(raw['steps'])]
    builtins = [{'id': move.id, 'name': move.variant('R')['name'], 'read_only': True,
                 'variants': [move.variant('R'), move.variant('L')]} for move in MOVES]
    return reference, builtins


def get_library(include_deleted=False):
    reference, builtins = reference_catalog()
    with _db() as conn:
        records = [json.loads(r['payload']) for r in conn.execute('SELECT payload FROM records ORDER BY rowid')]
    return {'schema_version': 1, 'reference': reference, 'builtins': builtins,
            'records': [r for r in records if include_deleted or not r['deleted']],
            'coverage': {'reference_records': len(reference), 'core_patterns': len(builtins),
                         'custom_records': sum(not r['deleted'] for r in records)}}


def get_record(record_id, version=None):
    with _db() as conn:
        return _get(conn, record_id, version)


def create_record(fields):
    fields = _fields(fields)
    with _db(True) as conn:
        return _put(conn, _new(fields))


def update_record(record_id, expected_version, fields):
    if not isinstance(fields, dict) or set(fields) - FIELDS:
        raise ValueError('Unsupported record fields.')
    with _db(True) as conn:
        record = _get(conn, record_id)
        _expect(record, expected_version)
        if record['deleted']:
            raise Conflict('This move is deleted. Its saved snapshots remain available.')
        changed = _next(record)
        changed.update(_fields({**{k: record[k] for k in FIELDS}, **fields}))
        changed['review'] = {'status': 'UNVERIFIED'}
        return _put(conn, changed)


def delete_record(record_id, expected_version):
    with _db(True) as conn:
        record = _get(conn, record_id)
        _expect(record, expected_version)
        changed = _next(record)
        changed['deleted'] = True
        return _put(conn, changed)


def restore_record(record_id, expected_version):
    """Undo soft deletion without overwriting any historical version."""
    with _db(True) as conn:
        record = _get(conn, record_id)
        _expect(record, expected_version)
        if not record['deleted']:
            return record
        changed = _next(record)
        changed['deleted'] = False
        return _put(conn, changed)


def reference_snapshot(reference_id, duration_counts=None):
    """Freeze reference text/facts as a manually usable, unverified definition.

Reference count facts may supply timing, but aggregate foot/turn metadata is
reported separately. It is not promoted to explicit per-event mechanics. A
caller can supply counts when using a variable-length pattern or annotation.
No automatic mirroring, generator admission, or source-rights claim occurs.
"""
    if not isinstance(reference_id, str) or not re.fullmatch(r'reference-(0|[1-9][0-9]{0,3})', reference_id):
        raise ValueError('Invalid reference identifier.')
    reference, _ = reference_catalog()
    index = int(reference_id.split('-')[1])
    if index >= len(reference):
        raise FileNotFoundError('Reference entry not found.')
    record = deepcopy(reference[index])
    fingerprint = hashlib.sha256(_json(record).encode()).hexdigest()
    duration = _count(duration_counts if duration_counts is not None else record.get('counts'), 'Duration', optional=True)
    definition_fingerprint = hashlib.sha256(_json({'reference': fingerprint, 'duration': duration}).encode()).hexdigest()
    event = {'text': record.get('description', ''), 'support_before': 'unknown', 'support_after': 'unknown'}
    result = {'id': reference_id, 'definition_id': reference_id,
              'snapshot_id': reference_id + '@' + definition_fingerprint[:16], 'reference_fingerprint': fingerprint,
              'name': record['name'], 'aliases': deepcopy(record['aliases']),
              'explanation': record.get('description', ''), 'events': [event],
              'review': {'status': 'UNVERIFIED'}, 'generator_eligible': False,
              'reference_record': record,
              'reported_mechanics': {k: record.get(k) for k in ('counts', 'count_notation', 'net_rotation_deg',
                  'foot_start', 'foot_end', 'weight_changes', 'travels', 'mirrorable')},
              'count_origin': 'USER_SUPPLIED' if duration_counts is not None else 'REFERENCE_REPORTED',
              'requires_counts': duration is None,
              'usage_note': 'Reference text and reported facts are retained. Choose the intended variant and review its explicit mechanics before validation or automatic generation.'}
    if duration is not None:
        result['duration_counts'] = duration
        event['duration_counts'] = duration
    return result


def snapshot(record_id, version=None):
    record = get_record(record_id, version)
    # Inline unknown event prevents a live catalog lookup or false implicit hold.
    events = deepcopy(record['events']) or [{'text': record['explanation'], 'support_before': 'unknown',
                                            'support_after': 'unknown'}]
    if not record['events'] and record['duration_counts'] is not None:
        events[0]['duration_counts'] = record['duration_counts']
    result = {'id': record['id'], 'definition_id': record['id'],
              'snapshot_id': record['id'] + '@' + str(record['version']),
              'name': record['name'], 'events': events, 'library_version': record['version'],
              'review': deepcopy(record['review']), 'explanation': record['explanation'],
              'attachments': deepcopy(record['attachments']), 'links': deepcopy(record['links']),
              'origins': deepcopy(record['origins'])}
    if record['duration_counts'] is not None:
        result['duration_counts'] = record['duration_counts']
    return result


def review_record(record_id, expected_version, reviewer, notes, confirmed):
    if confirmed is not True:
        raise ValueError('Explicit review confirmation is required.')
    reviewer = _text(reviewer, 'Reviewer', 200, True)
    notes = _text(notes, 'Review notes', 3000)
    with _db(True) as conn:
        record = _get(conn, record_id)
        _expect(record, expected_version)
        if record['deleted']:
            raise Conflict('Deleted moves cannot be reviewed.')
        support = record['events'][0].get('support_before', 'unknown') if record['events'] else 'unknown'
        document = {'start': {'support': 'L' if support == 'any' else support, 'facing_deg': '0'},
                    'parts': [{'id': 'A', 'moves': [{k: record[k] for k in ('id', 'name', 'duration_counts', 'events')}]}],
                    'routine': [{'id': 'first', 'part_id': 'A'}]}
        report = compile_choreography(document)
        if not report['verified']:
            raise ValueError('Complete and reconcile the explicit mechanics before author review: ' +
                             '; '.join(i.get('message', i.get('code', '')) for i in report['issues'][:4]))
        changed = _next(record)
        changed['review'] = {'status': 'AUTHOR_REVIEWED', 'reviewer': reviewer, 'notes': notes, 'at': _now()}
        return _put(conn, changed)


def _filename(value):
    value = _text(value, 'Filename', 180, True)
    if any(c in value for c in '/\\:') or value in {'.', '..'} or any(ord(c) < 32 for c in value):
        raise ValueError('Use a plain filename without directories.')
    return value


def _table_rows(headers, rows):
    headers = [str(h).strip() if h is not None else '' for h in headers]
    if not headers or len(headers) > 64 or any(not h for h in headers) or len(set(headers)) != len(headers):
        raise ValueError('Columns must have distinct nonempty names.')
    result = []
    for number, row in enumerate(rows, 1):
        if number > MAX_ROWS:
            raise ValueError('Import has more than 1000 rows, including blank rows.')
        if not any(v not in (None, '') for v in row):
            continue
        if len(row) > len(headers):
            raise ValueError('A row has more values than column names.')
        result.append(dict(zip(headers, list(row) + [''] * (len(headers) - len(row)))))
        if len(result) > MAX_ROWS:
            raise ValueError('Import has more than 1000 rows.')
    return result


def _parse_import(filename, payload):
    filename = _filename(filename)
    if not payload or len(payload) > MAX_IMPORT_BYTES:
        raise ValueError('Import must be nonempty and at most 8 MiB.')
    suffix = Path(filename).suffix.casefold()
    try:
        if suffix == '.json':
            rows = json.loads(payload.decode('utf-8-sig'), parse_constant=lambda _: (_ for _ in ()).throw(ValueError('Nonfinite JSON values are not allowed.')))
            if isinstance(rows, dict):
                _no_secrets(rows)
                rows = rows.get('moves', rows.get('records'))
            if not isinstance(rows, list) or any(not isinstance(row, dict) for row in rows):
                raise ValueError('JSON must contain a list of move objects, or a moves/records list.')
        elif suffix == '.csv':
            reader = csv.reader(io.StringIO(payload.decode('utf-8-sig'), newline=''), strict=True)
            headers = next(reader, [])
            rows = _table_rows(headers, reader)
        elif suffix == '.xlsx':
            from openpyxl import load_workbook
            with zipfile.ZipFile(io.BytesIO(payload)) as archive:
                entries = archive.infolist()
                if len(entries) > 2000 or sum(i.file_size for i in entries) > MAX_ARCHIVE_BYTES:
                    raise ValueError('Spreadsheet archive is too large.')
                for info in entries:
                    path = PurePosixPath(info.filename.replace('\\', '/'))
                    if path.is_absolute() or '..' in path.parts or ':' in info.filename or info.flag_bits & 1:
                        raise ValueError('Unsupported spreadsheet archive entry.')
            workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=False, keep_links=False)
            try:
                sheet = workbook.worksheets[0]
                if (sheet.max_row or 0) > MAX_ROWS + 1 or (sheet.max_column or 0) > 64:
                    raise ValueError('Spreadsheet must have at most 1000 data rows and 64 columns.')
                def values():
                    for row in sheet.iter_rows():
                        if any(cell.data_type == 'f' for cell in row):
                            raise ValueError('Replace spreadsheet formulas with values before importing.')
                        yield [cell.value for cell in row]
                iterator = values()
                rows = _table_rows(next(iterator, []), iterator)
            finally:
                workbook.close()
        else:
            raise ValueError('Supported import formats are CSV, XLSX and JSON.')
    except (UnicodeError, json.JSONDecodeError, csv.Error, zipfile.BadZipFile, KeyError, IndexError) as exc:
        raise ValueError('The import file is malformed or unreadable.') from exc
    if not rows or len(rows) > MAX_ROWS:
        raise ValueError('Import must contain 1 to 1000 move rows.')
    _no_secrets(rows)
    return rows


def _names(record):
    return {str(n).strip().casefold() for n in [record.get('name', ''), *record.get('aliases', [])] if str(n).strip()}


def _duplicate_index(conn):
    reference, builtins = reference_catalog()
    custom = [json.loads(row['payload']) for row in conn.execute('SELECT payload FROM records')]
    return [('reference', row) for row in reference] + [('core', row) for row in builtins] + [('user', row) for row in custom if not row['deleted']]


def _duplicates(fields, index):
    names = _names(fields)
    return [{'kind': kind, 'id': r['id'], 'name': r['name'], **({'version': r['version']} if kind == 'user' else {})}
            for kind, r in index if _names(r) & names]


def preview_import(filename, payload, field_mapping=None, origin_kind='file'):
    if origin_kind not in {'file', 'ai_user_supplied'}:
        raise ValueError('Choose file or ai_user_supplied origin.')
    field_mapping = field_mapping or {}
    if not isinstance(field_mapping, dict) or any(not isinstance(k, str) or v not in FIELDS for k, v in field_mapping.items()):
        raise ValueError('Map source column names to supported move fields.')
    if len(set(field_mapping.values())) != len(field_mapping):
        raise ValueError('Each target field may be mapped once.')
    rows = _parse_import(filename, payload)
    with _db(True) as conn:
        index = _duplicate_index(conn)
        preview_rows, prior_names = [], []
        for number, raw in enumerate(rows, 1):
            fields, errors, warnings = {}, [], []
            if any(k in raw for k in ('review', 'generator_eligible', 'in_generator', 'verified')):
                warnings.append('Imported review and generator claims are discarded. You must review the move locally.')
            if Path(filename).suffix.casefold() == '.xlsx':
                warnings.append('Only the first worksheet is included.')
            try:
                targets = [field_mapping.get(k, k) for k in raw if field_mapping.get(k, k) in FIELDS]
                if len(set(targets)) != len(targets):
                    raise ValueError('Multiple source columns map to the same field. Rename or remove the conflicting column.')
                mapped = {field_mapping.get(k, k): value for k, value in raw.items() if field_mapping.get(k, k) in FIELDS}
                for field in ('events', 'origins', 'links'):
                    if isinstance(mapped.get(field), str):
                        mapped[field] = json.loads(mapped[field]) if mapped[field].strip() else []
                fields = _fields(mapped)
            except (ValueError, TypeError, ZeroDivisionError) as exc:
                errors.append(str(exc))
            names = _names(fields)
            duplicates = _duplicates(fields, index) if fields else []
            duplicate_rows = [n for n, prior in prior_names if names & prior]
            prior_names.append((number, names))
            preview_rows.append({'row': number, 'fields': fields, 'errors': errors, 'warnings': warnings,
                                 'duplicates': duplicates, 'duplicate_rows': duplicate_rows})
        preview = {'schema_version': 1, 'token': _id('preview'), 'filename': _filename(filename),
                   'input_sha256': hashlib.sha256(payload).hexdigest(), 'origin_kind': origin_kind,
                   'field_mapping': field_mapping, 'rows': preview_rows, 'created_at': _now()}
        digest = hashlib.sha256(_json(preview).encode()).hexdigest()
        conn.execute('INSERT INTO previews VALUES(?,?,?,?,NULL)', (preview['token'], digest, _json(preview), preview['created_at']))
        return {**preview, 'digest': digest}


def commit_import(token, digest, decisions, confirmed):
    _check_id(token, 'preview')
    if confirmed is not True:
        raise ValueError('Explicit approval of the preview is required.')
    if not isinstance(decisions, list) or not decisions or len(decisions) > MAX_ROWS:
        raise ValueError('Select rows to create, update, keep as variants, or skip.')
    _no_secrets(decisions)
    decision_digest = hashlib.sha256(_json(decisions).encode()).hexdigest()
    with _db(True) as conn:
        stored = conn.execute('SELECT * FROM previews WHERE id=?', (token,)).fetchone()
        if stored is None:
            raise FileNotFoundError('Import preview not found.')
        if digest != stored['digest'] or hashlib.sha256(stored['payload'].encode()).hexdigest() != stored['digest']:
            raise Conflict('The preview changed. Preview this file again.')
        if stored['result']:
            result = json.loads(stored['result'])
            if result['decision_digest'] != decision_digest:
                raise Conflict('This preview was already committed with different decisions.')
            return result
        if datetime.fromisoformat(stored['created']) < datetime.now(timezone.utc) - timedelta(days=7):
            raise Conflict('This preview expired. Preview this file again.')
        preview = json.loads(stored['payload'])
        rows = {r['row']: r for r in preview['rows']}
        index = _duplicate_index(conn)
        changes, records, seen_rows, seen_targets = [], [], set(), set()
        for decision in decisions:
            if not isinstance(decision, dict) or set(decision) - {'row', 'action', 'target_id', 'expected_version'}:
                raise ValueError('Invalid import decision fields.')
            number, action = decision.get('row'), decision.get('action')
            if type(number) is not int or number not in rows or number in seen_rows:
                raise ValueError('Each decision must identify one distinct preview row.')
            seen_rows.add(number)
            if action == 'skip':
                continue
            row = rows[number]
            if row['errors']:
                raise ValueError('Row ' + str(number) + ' has errors. Correct the source and preview again.')
            fields = deepcopy(row['fields'])
            fields['origins'].append({'kind': preview['origin_kind'], 'label': preview['filename'],
                                      'locator': 'row ' + str(number),
                                      'rights_note': 'User supplied; rights and mechanics have not been independently checked.'})
            prior = None
            if action == 'update':
                prior = _get(conn, decision.get('target_id'))
                _expect(prior, decision.get('expected_version'))
                if prior['deleted'] or prior['id'] in seen_targets:
                    raise Conflict('An update target is deleted or selected more than once.')
                seen_targets.add(prior['id'])
                record = _next(prior)
                record.update(fields)
                record['review'] = {'status': 'UNVERIFIED'}
            elif action in {'create', 'variant'}:
                if action == 'create' and _duplicates(fields, index):
                    raise Conflict('Row ' + str(number) + ' matches another move. Explicitly keep it as a variant or update a user record.')
                record = _new(fields)
            else:
                raise ValueError('Choose create, variant, update or skip.')
            # Detect conflicts against earlier selected rows without partial writes.
            index = [(kind, r) for kind, r in index if r['id'] != record['id']]
            index.append(('user', record))
            changes.append({'id': record['id'], 'applied_version': record['version'], 'prior': prior})
            records.append(record)
        batch_id = _id('batch')
        for record in records:
            _put(conn, record)
        result = {'batch_id': batch_id, 'records': records, 'applied': len(records),
                  'skipped': len(preview['rows']) - len(records), 'decision_digest': decision_digest}
        conn.execute('INSERT INTO batches VALUES(?,?,?)', (batch_id, _json(changes), 'applied'))
        conn.execute('UPDATE previews SET result=? WHERE id=?', (_json(result), token))
        return result


def rollback_batch(batch_id, confirmed):
    _check_id(batch_id, 'batch')
    if confirmed is not True:
        raise ValueError('Explicit rollback confirmation is required.')
    with _db(True) as conn:
        batch = conn.execute('SELECT * FROM batches WHERE id=?', (batch_id,)).fetchone()
        if batch is None:
            raise FileNotFoundError('Import batch not found.')
        if batch['state'] == 'rolled_back':
            return {'batch_id': batch_id, 'state': 'rolled_back'}
        changes = json.loads(batch['changes'])
        current = [_get(conn, change['id']) for change in changes]
        for record, change in zip(current, changes):
            _expect(record, change['applied_version'])
        for record, change in zip(current, changes):
            restored = _next(record)
            if change['prior'] is None:
                restored['deleted'] = True
            else:
                restored = {**deepcopy(change['prior']), 'version': restored['version'], 'updated_at': restored['updated_at']}
            _put(conn, restored)
        conn.execute('UPDATE batches SET state=? WHERE id=?', ('rolled_back', batch_id))
        return {'batch_id': batch_id, 'state': 'rolled_back', 'records': len(changes)}


def template(format_name):
    headers = ['name', 'aliases', 'explanation', 'duration_counts', 'difficulty']
    example = ['My custom step', 'My alias', 'Replace this with your own instructions.', '3/2', 'Unspecified']
    if format_name == 'json':
        return (json.dumps([dict(zip(headers, example))], indent=2).encode(), 'application/json')
    if format_name == 'csv':
        output = io.StringIO(newline='')
        writer = csv.writer(output)
        writer.writerow(headers)
        writer.writerow(example)
        return output.getvalue().encode('utf-8-sig'), 'text/csv; charset=utf-8'
    if format_name == 'xlsx':
        from openpyxl import Workbook
        workbook = Workbook()
        workbook.active.append(headers)
        workbook.active.append(example)
        output = io.BytesIO()
        workbook.save(output)
        workbook.close()
        return output.getvalue(), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    raise ValueError('Choose json, csv or xlsx.')


def _media_type(filename, payload):
    suffix = Path(_filename(filename)).suffix.casefold()
    images = {'.png': ('PNG', 'image/png'), '.jpg': ('JPEG', 'image/jpeg'),
              '.jpeg': ('JPEG', 'image/jpeg'), '.webp': ('WEBP', 'image/webp'), '.gif': ('GIF', 'image/gif')}
    videos = {'.mp4': 'video/mp4', '.mov': 'video/quicktime', '.webm': 'video/webm'}
    limit = MAX_IMAGE_BYTES if suffix in images else MAX_VIDEO_BYTES
    if not payload or len(payload) > limit:
        raise ValueError('Media is empty or exceeds the image (20 MiB) / video (100 MiB) limit.')
    if suffix in images:
        from PIL import Image, UnidentifiedImageError
        try:
            with Image.open(io.BytesIO(payload)) as image:
                if image.format != images[suffix][0] or image.width * image.height > 40000000:
                    raise ValueError('Image format or dimensions are unsupported.')
                image.verify()
        except (OSError, SyntaxError, UnidentifiedImageError, Image.DecompressionBombError) as exc:
            raise ValueError('This image is invalid or too large to decode safely.') from exc
        return suffix, images[suffix][1]
    if suffix in videos:
        valid = (len(payload) >= 12 and payload[4:8] == b'ftyp') if suffix in {'.mp4', '.mov'} else payload.startswith(b'\x1a\x45\xdf\xa3')
        if not valid:
            raise ValueError('Video header does not match the selected format.')
        return suffix, videos[suffix]
    raise ValueError('Supported media: PNG, JPEG, WebP, GIF, MP4, MOV and WebM.')


def attach_media(record_id, expected_version, filename, payload, caption='', orientation='unspecified',
                 start_seconds=None, end_seconds=None):
    filename = _filename(filename)
    suffix, mime = _media_type(filename, payload)
    caption = _text(caption, 'Caption', 2000)
    if orientation not in {'unspecified', 'front', 'back', 'left', 'right', 'overhead'}:
        raise ValueError('Choose a supported camera orientation.')
    start = _count(start_seconds, 'Range start', optional=True)
    end = _count(end_seconds, 'Range end', optional=True)
    if (start is None) != (end is None) or (start is not None and exact_count(end) <= exact_count(start)):
        raise ValueError('Supply both range endpoints with end after start.')
    if mime.startswith('image/') and start is not None:
        raise ValueError('Time ranges apply to video attachments.')
    media_id = _id('media')
    metadata = {'id': media_id, 'original_filename': filename, 'mime': mime, 'bytes': len(payload),
                'sha256': hashlib.sha256(payload).hexdigest(), 'caption': caption, 'orientation': orientation,
                'start_seconds': start, 'end_seconds': end,
                'range_review': 'USER_ENTERED_NOT_MEASURED' if start is not None else 'NOT_SPECIFIED'}
    media_root = directory() / 'media'
    media_root.mkdir(parents=True, exist_ok=True)
    destination = media_root / (media_id + suffix)
    temporary = media_root / (media_id + '.tmp')
    try:
        with _db(True) as conn:
            record = _get(conn, record_id)
            _expect(record, expected_version)
            if record['deleted']:
                raise Conflict('Deleted moves cannot receive attachments.')
            with temporary.open('xb') as handle:
                handle.write(payload)
                handle.flush()
                os.fsync(handle.fileno())
            os.replace(temporary, destination)
            conn.execute('INSERT INTO media VALUES(?,?,?,?)', (media_id, record_id, destination.name, _json(metadata)))
            changed = _next(record)
            changed['attachments'].append(metadata)
            _put(conn, changed)
        return {'record': changed, 'attachment': metadata}
    except BaseException:
        temporary.unlink(missing_ok=True)
        destination.unlink(missing_ok=True)
        raise


def media_content(media_id):
    _check_id(media_id, 'media')
    with _db() as conn:
        row = conn.execute('SELECT * FROM media WHERE id=?', (media_id,)).fetchone()
        if row is None:
            raise FileNotFoundError('Attachment not found.')
        metadata = json.loads(row['payload'])
        root = (directory() / 'media').resolve()
        path = (root / row['filename']).resolve()
        if path.parent != root or not path.is_file():
            raise FileNotFoundError('Attachment file is missing. The move and its saved versions are retained.')
        return path, metadata
