"""Sheets and portable drafts from the shared exact choreography timeline.

Native packages contain a saved working draft and frozen movement events, not
audio files, provider settings, local file paths, or project history. Referenced
media is described for explicit relinking. Import never extracts archive paths.
"""
from copy import deepcopy
import csv
from fractions import Fraction
import hashlib
from html import escape
from io import BytesIO, StringIO
import json
import math
from pathlib import PurePosixPath
import re
import stat
from urllib.parse import urlsplit
import zipfile

from .choreography import compile_choreography, exact_count
from .music_map import count_to_seconds, validate_map


PACKAGE_KIND = "line-dance-project"
PACKAGE_SCHEMA = 1
MAX_INPUT_BYTES = 16 * 1024 * 1024
MAX_MEMBER_BYTES = 12 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 16 * 1024 * 1024
MAX_MEMBERS = 16
MAX_COMPRESSION_RATIO = 200
FORMATS = {"pdf", "docx", "txt", "html", "csv", "xlsx", "srt", "vtt", "json", "zip"}
META_FIELDS = {"dance_title", "choreographer", "country", "contact", "level_label", "description",
               "signature_note", "ending_note", "youtube_url", "sheet_url", "song_url", "spotify_url", "source_url", "credits", "links", "print_qr"}
MUSIC_FIELDS = {"bpm", "first_count", "meter", "key", "anchors", "measured_bpm", "manually_corrected", "timing_confirmed"}
PRIVATE_FIELDS = {"api_key", "apikey", "api_key_protected", "access_token", "refresh_token", "authorization", "password", "secret",
                  "client_secret", "client_id", "api_token", "credentials", "provider_settings", "ai_settings", "connection", "endpoint", "path", "local_path", "file_path", "source_path", "audio_data", "base64", "data_url"}


class ExportError(ValueError):
    pass


def _json_bytes(value):
    return (json.dumps(value, ensure_ascii=False, sort_keys=True, indent=2, allow_nan=False) + "\n").encode("utf-8")


def _text(value, limit=20000):
    if value is None:
        return ""
    if isinstance(value, list):
        value = "; ".join(_text(item) for item in value)
    if not isinstance(value, (str, int, float, bool)):
        return ""
    value = str(value)
    if len(value) > limit:
        raise ExportError("A text field is too long to export. Shorten it or split the instructions.")
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", value)


def safe_url(value):
    value = _text(value, 2048).strip()
    if not value or any(ord(char) < 32 for char in value):
        return None
    try:
        parsed = urlsplit(value)
        if parsed.scheme.lower() not in ("http", "https") or not parsed.netloc or parsed.username or parsed.password:
            return None
        _ = parsed.port
    except ValueError:
        return None
    return value


def _redact(value, depth=0):
    if depth > 60:
        raise ExportError("Project nesting exceeds the portable format limit.")
    if isinstance(value, dict):
        return {str(key): _redact(child, depth + 1) for key, child in value.items()
                if re.sub(r"[- ]", "_", str(key).lower()) not in PRIVATE_FIELDS}
    if isinstance(value, list):
        return [_redact(child, depth + 1) for child in value]
    if isinstance(value, float) and not math.isfinite(value):
        raise ExportError("Project contains a nonfinite numeric value.")
    return deepcopy(value)


def _metadata(project):
    draft = project.get("draft") or {}
    if not isinstance(draft, dict):
        raise ExportError("The saved working draft is invalid.")
    meta = draft.get("sheet_meta", project.get("sheet_meta", {})) or {}
    song = project.get("song") or draft.get("song") or {}
    if not isinstance(meta, dict) or not isinstance(song, dict):
        raise ExportError("Song and sheet details must be objects.")
    return draft, meta, song


def _compile_project(project):
    draft, _, _ = _metadata(project)
    choreography = draft.get("choreography")
    if choreography is None:
        from . import project as store
        editor = draft.get("editor") or {}
        items = editor.get("moves", [])
        try:
            choreography = store.resolve_move_variants(project, items, draft=True)
        except ValueError as exc:
            raise ExportError(str(exc)) from exc
    snapshots = draft.get("movement_snapshots") or project.get("movement_snapshots") or {}
    try:
        return compile_choreography(choreography, snapshots)
    except (TypeError, ValueError, KeyError, RecursionError) as exc:
        raise ExportError("The choreography cannot be compiled; review its saved structure.") from exc


def _public_spotify(meta):
    """Share link only. A missing or unusable value is omitted from the sheet."""
    from .song_card import normalize_spotify_url
    try:
        return normalize_spotify_url(meta.get("spotify_url")) or None
    except ValueError:
        return None


def _links(meta):
    values = [("Video demonstration", meta.get("youtube_url")), ("Step sheet", meta.get("sheet_url")),
              ("Spotify link", _public_spotify(meta)), ("Music", meta.get("song_url")),
              ("Movement source", meta.get("source_url"))]
    for entry in meta.get("links", []) if isinstance(meta.get("links"), list) else []:
        if isinstance(entry, str):
            entry = {"url": entry}
        if isinstance(entry, dict):
            values.append((_text(entry.get("label")) or "Reference", entry.get("url")))
    out, seen = [], set()
    for label, raw in values:
        url = safe_url(raw)
        if url and url not in seen:
            seen.add(url)
            out.append({"label": _text(label), "url": url})
    return out


def _annotation_fields(raw, fields):
    if not isinstance(raw, dict):
        return {}
    return {key: deepcopy(value) for key, value in raw.items() if key in fields
            and (value is None or isinstance(value, (str, bool, int, float))
                 or isinstance(value, list) and all(isinstance(item, str) for item in value))}


def _phrase_origin(raw, depth=0):
    """Portable phrase provenance, with bounded history and no arbitrary fields."""
    if depth > 16:
        raise ExportError("Phrase ancestry exceeds 16 generations. Use a full private backup to preserve this project.")
    if not isinstance(raw, dict):
        raise ExportError("Phrase provenance must be an object.")
    result = {}
    for key in ("phrase_id", "source_occurrence_id", "source_snapshot_id"):
        value = raw.get(key)
        if value is None:
            if key in raw:
                result[key] = None
        elif isinstance(value, str) and re.fullmatch(r"[A-Za-z0-9_.:@+\-]{1,200}", value):
            result[key] = value
        else:
            raise ExportError("Phrase provenance contains an invalid " + key + ".")
    for key in ("source_definition_hash", "transformed_hash"):
        value = raw.get(key)
        if value is None:
            if key in raw:
                result[key] = None
        elif isinstance(value, str) and re.fullmatch(r"[a-fA-F0-9]{64}", value):
            result[key] = value
        else:
            raise ExportError("Phrase provenance contains an invalid definition fingerprint.")
    for key in ("name", "status", "original_review_status", "transform"):
        value = raw.get(key)
        if value is None:
            if key in raw:
                result[key] = None
            continue
        if not isinstance(value, str) or len(value) > 400:
            raise ExportError("Phrase provenance text is malformed or too long.")
        if re.match(r"^[A-Za-z]:[\\/]", value) or value.startswith(("/", "\\")):
            value = re.split(r"[\\/]", value)[-1]
        result[key] = value
    if "version" in raw:
        if type(raw["version"]) is not int or not 1 <= raw["version"] <= 1000000000:
            raise ExportError("Phrase provenance has an invalid version.")
        result["version"] = raw["version"]
    if "mirrored" in raw:
        if type(raw["mirrored"]) is not bool:
            raise ExportError("Phrase provenance has an invalid mirror flag.")
        result["mirrored"] = raw["mirrored"]
    if "prior_phrase_origin" in raw:
        prior = raw["prior_phrase_origin"]
        result["prior_phrase_origin"] = None if prior is None else _phrase_origin(prior, depth + 1)
    return result


def _annotations(raw):
    """Author/source notes are portable data, never substitute mechanics."""
    result = _annotation_fields(raw, {"explanation", "description", "notes", "annotation", "author", "choreographer", "credits",
                                     "aliases", "difficulty", "library_version", "snapshot_id", "reference_fingerprint",
                                     "count_origin", "usage_note", "locked"})
    if not isinstance(raw, dict):
        return result
    # Preserve the frozen definition's identity/review provenance independently
    # of the normalized compiler graph. These fields remain descriptive data;
    # importing a claim never admits the move to an automatic generator.
    for key in ("source_hash", "definition_hash"):
        if isinstance(raw.get(key), str) and re.fullmatch(r"[a-fA-F0-9]{64}", raw[key]):
            result[key] = raw[key]
    for key in ("definition_id", "move_id", "source_pack"):
        if isinstance(raw.get(key), str) and re.fullmatch(r"[A-Za-z0-9_.:@+\-]{1,160}", raw[key]):
            result[key] = raw[key]
    result.update(_annotation_fields(raw, {"group", "family", "level", "review_status", "review_note", "provenance_note",
                                           "required_start_facing", "start_free_foot", "net_rotation_deg", "travel"}))
    for key in ("generator_eligible", "in_generator", "manual_only", "mechanically_complete", "ab_safe"):
        if type(raw.get(key)) is bool:
            result[key] = raw[key]
    if isinstance(raw.get("source_ids"), list):
        result["source_ids"] = [sid for sid in raw["source_ids"]
                                if isinstance(sid, str) and re.fullmatch(r"[A-Za-z0-9_.:@+\-]{1,160}", sid)]
    if isinstance(raw.get("sources"), list):
        result["sources"] = []
        for source in raw["sources"][:100]:
            if not isinstance(source, dict):
                continue
            url = safe_url(source.get("url"))
            if not url:
                continue
            item = {"title": _text(source.get("title"), 1000), "url": url}
            if isinstance(source.get("id"), str) and re.fullmatch(r"[A-Za-z0-9_.:@+\-]{1,160}", source["id"]):
                item["id"] = source["id"]
            result["sources"].append(item)
    result["links"] = _links(raw)
    result["origins"] = []
    for origin in raw.get("origins", []) if isinstance(raw.get("origins"), list) else []:
        item = _annotation_fields(origin, {"kind", "label", "url", "locator", "rights_note"})
        if "url" in item:
            item["url"] = safe_url(item["url"]) or ""
        for key in ("label", "locator"):
            value = item.get(key)
            if isinstance(value, str) and (re.match(r"^[A-Za-z]:[\\/]", value) or value.startswith(("/", "\\"))):
                item[key] = re.split(r"[\\/]", value)[-1]
        result["origins"].append(item)
    if isinstance(raw.get("review"), dict):
        result["review"] = _annotation_fields(raw["review"], {"status", "reviewer", "notes", "at"})
    if isinstance(raw.get("reference_record"), dict):
        result["reference_record"] = _annotation_fields(raw["reference_record"], {
            "id", "name", "aliases", "description", "notation", "notes", "counts", "count_notation", "net_rotation_deg",
            "foot_start", "foot_end", "weight_changes", "travels", "mirrorable", "level", "flags", "ab_safe", "seed_verified"})
    if isinstance(raw.get("reported_mechanics"), dict):
        result["reported_mechanics"] = _annotation_fields(raw["reported_mechanics"], {
            "counts", "count_notation", "net_rotation_deg", "foot_start", "foot_end", "weight_changes", "travels", "mirrorable"})
    if isinstance(raw.get("attachments"), list):
        result["attachments"] = _media_references({"draft": {"attachments": raw["attachments"]}})
    if "phrase_origin" in raw:
        result["phrase_origin"] = _phrase_origin(raw["phrase_origin"])
    return result


def _annotated_document(project, compiled):
    draft = project.get("draft") or {}
    source = draft.get("choreography")
    normalized = deepcopy(compiled["normalized_document"])
    if not isinstance(source, (dict, list)):
        return normalized
    sources = [{"moves": source}] if isinstance(source, list) else source.get("parts", [])
    if not isinstance(sources, list):
        return normalized
    snapshots = draft.get("movement_snapshots") or project.get("movement_snapshots") or {}
    def enrich(target, raw):
        if not isinstance(raw, dict):
            return
        origin = snapshots.get(str(raw.get("snapshot_id")), {}) if isinstance(snapshots, dict) else {}
        if not isinstance(origin, dict):
            origin = {}
        if isinstance(raw.get("definition"), dict):
            origin = {**origin, **raw["definition"]}
        target.update(_annotations({**origin, **raw}))
    for part, raw_part in zip(normalized["parts"], sources):
        if not isinstance(raw_part, dict) or not isinstance(raw_part.get("moves", []), list):
            continue
        enrich(part, raw_part)
        for move, raw_move in zip(part["moves"], raw_part.get("moves", [])):
            enrich(move, raw_move)
    raw_runs = source.get("routine", []) if isinstance(source, dict) else []
    if not isinstance(raw_runs, list):
        raw_runs = []
    for run, raw_run in zip(normalized["routine"], raw_runs):
        if not isinstance(raw_run, dict):
            continue
        enrich(run, raw_run)
        for mid, move in run.get("overrides", {}).items():
            enrich(move, (raw_run.get("overrides") or {}).get(mid, {}))
    return normalized


def _used_annotations(project, compiled):
    document = _annotated_document(project, compiled)
    parts = {part["id"]: part for part in document["parts"]}
    runs = {run["id"]: run for run in document["routine"]}
    used_by_occurrence = {}
    for event in compiled["events"]:
        used_by_occurrence.setdefault(event["occurrence_id"], set()).add(event["move_id"])
    for occurrence in compiled["occurrences"]:
        part = parts[occurrence["part_id"]]
        yield part
        used = used_by_occurrence.get(occurrence["id"], set())
        overrides = runs[occurrence["routine_id"]].get("overrides") or {}
        for move in part["moves"]:
            if move["id"] in used:
                yield overrides.get(move["id"], move)


def build_sheet_model(project):
    """One compiled presentation model for every sheet/caption renderer."""
    draft, meta, song = _metadata(project)
    compiled = _compile_project(project)
    from .project_media import recording_review_needed
    recording_pending = recording_review_needed(project)
    occurrences = []
    by_id = {event["id"]: event for event in compiled["events"]}
    for occurrence in compiled["occurrences"]:
        rows = []
        for eid in occurrence["event_ids"]:
            event = by_id[eid]
            lines = event.get("lines") or []
            instructions = " ".join(_text(line.get("text")) for line in lines if isinstance(line, dict)) or _text(event.get("text"))
            rows.append({"id": eid, "start_count": event["start_count"], "end_count": event["end_count"],
                         "duration_counts": event["duration_counts"], "label": event["count_label"],
                         "move": _text(event.get("move_name")), "instructions": instructions,
                         "notes": _text(event.get("notes")), "facing": event["state_after"].get("facing_deg"),
                         "support": event["state_after"].get("support", "unknown"),
                         "part_id": occurrence["part_id"], "occurrence_id": occurrence["id"]})
        occurrences.append({**deepcopy(occurrence), "rows": rows})
    music_map = draft.get("music_map") or {}
    if not isinstance(music_map, dict):
        raise ExportError("The music timing map must be an object.")
    links = _links(meta)
    known_urls = {link["url"] for link in links}
    for item in _used_annotations(project, compiled):
        references = list(item.get("links") or [])
        references += [{"label": source.get("title") or "Movement source", "url": source.get("url")} for source in item.get("sources", [])]
        references += [{"label": origin.get("label") or "Movement source", "url": origin.get("url")} for origin in item.get("origins", [])]
        references += [{"label": media.get("caption") or "Movement media", "url": media.get("url")} for media in item.get("attachments", [])]
        for link in references:
            url = safe_url(link.get("url"))
            if url and url not in known_urls:
                known_urls.add(url)
                links.append({"label": _text(link.get("label")) or _text(item.get("name")) or "Movement reference", "url": url})
    return {"title": _text(meta.get("dance_title") or project.get("name") or "Untitled dance"),
            "choreographer": _text(meta.get("choreographer")) or "Not specified",
            "country": _text(meta.get("country")), "contact": _text(meta.get("contact")),
            "level": _text(meta.get("level_label")) or "Not specified",
            "song_title": _text(song.get("title")), "artist": _text(song.get("artist")),
            "credits": _text(meta.get("credits")), "notes": _text(meta.get("description")),
            "signature_note": _text(meta.get("signature_note")), "ending_note": _text(meta.get("ending_note")),
            "links": links, "music_map": deepcopy(music_map),
            "compiled": compiled, "occurrences": occurrences,
            "status": ("RECORDING TIMING NEEDS REVIEW · " if recording_pending else "") + ("DRAFT - FAILED CHECKS" if not compiled["valid"] else "DRAFT - UNVERIFIED MECHANICS" if not compiled["verified"] else "Working draft - modeled checks pass; instructor review required")}


def _summary(model):
    compiled = model["compiled"]
    items = [f"Choreographer: {model['choreographer']}" + (f" ({model['country']})" if model["country"] else ""),
             f"Counts: {compiled['total_counts']} | Level: {model['level']} | Grouping: {compiled['meter']['group_counts']}",
             f"Music: {model['song_title'] or 'Not selected'}" + (f" - {model['artist']}" if model["artist"] else ""),
             f"Start: {compiled['start_state']['free_foot']} free; facing {compiled['start_state']['facing_deg'] or 'unknown'} degrees",
             model["status"]]
    for label, key in (("Contact", "contact"), ("Credits", "credits"), ("Notes", "notes"), ("Signature", "signature_note"), ("Ending", "ending_note")):
        if model[key]:
            items.append(f"{label}: {model[key]}")
    return items


def _occurrence_title(occurrence):
    value = f"{occurrence['part_name']} | {occurrence['kind']} | occurrence {occurrence['occurrence_number']} | {occurrence['duration_counts']} counts"
    if occurrence.get("reason"):
        value += f" | {occurrence['reason']} after {occurrence['duration_counts']} counts"
    return value


def _row_text(row):
    value = row["instructions"] or row["move"]
    if row["notes"]:
        value += " (" + row["notes"] + ")"
    return value


def _checks(model):
    return [f"{issue['severity']}: {issue['message']} {issue['action']}" for issue in model["compiled"]["issues"]]


def render_txt(model):
    out = [model["title"], "", *_summary(model), ""]
    for occurrence in model["occurrences"]:
        out.append(_occurrence_title(occurrence))
        out.extend(f"{row['label']} [{row['duration_counts']} counts]  {_row_text(row)}" for row in occurrence["rows"])
        out.append("")
    if model["links"]:
        out.extend(["References", *[f"{link['label']}: {link['url']}" for link in model["links"]], ""])
    if _checks(model):
        out.extend(["Checks and review", *_checks(model)])
    return ("\n".join(out).rstrip() + "\n").encode("utf-8")


def _qr_png(url):
    from reportlab.graphics.barcode.qr import QrCodeWidget
    from PIL import Image, ImageDraw
    qr = QrCodeWidget(url)
    qr.qr.make()
    matrix = qr.qr.modules
    scale, quiet = 6, 4
    image = Image.new("RGB", ((len(matrix) + quiet * 2) * scale,) * 2, "white")
    draw = ImageDraw.Draw(image)
    for y, row in enumerate(matrix):
        for x, dark in enumerate(row):
            if dark:
                draw.rectangle(((x + quiet) * scale, (y + quiet) * scale, (x + quiet + 1) * scale - 1, (y + quiet + 1) * scale - 1), fill="black")
    output = BytesIO()
    image.save(output, format="PNG")
    return output.getvalue()


def render_html(model, paper="letter", large_print=False, print_qr=False):
    import base64
    size = "16pt" if large_print else "11pt"
    out = ["<!doctype html><html lang='en'><meta charset='utf-8'>", f"<title>{escape(model['title'])}</title>",
           f"<style>@page{{size:{paper};margin:18mm}}body{{font:{size}/1.45 Arial,sans-serif;color:#111;max-width:900px;margin:24px auto;padding:0 16px}}h1{{font-size:1.8em}}h2{{font-size:1.15em;break-after:avoid}}table{{width:100%;border-collapse:collapse;margin:12px 0 24px}}th,td{{text-align:left;vertical-align:top;border-bottom:1px solid #ccc;padding:7px}}th:first-child,td:first-child{{width:100px}}a{{color:#174783;overflow-wrap:anywhere}}.qr{{width:90px;height:90px}}.notes{{white-space:pre-wrap}}@media print{{body{{margin:0;padding:0;max-width:none}}}}</style>",
           "<body>", f"<h1>{escape(model['title'])}</h1>", *[f"<p class='notes'>{escape(line)}</p>" for line in _summary(model)]]
    for occurrence in model["occurrences"]:
        out += [f"<h2>{escape(_occurrence_title(occurrence))}</h2>", "<table><thead><tr><th>Count</th><th>Instructions</th></tr></thead><tbody>"]
        for row in occurrence["rows"]:
            out.append(f"<tr><td>{escape(row['label'])}<br><small>{escape(row['duration_counts'])} counts</small></td><td class='notes'>{escape(_row_text(row))}</td></tr>")
        out.append("</tbody></table>")
    if model["links"]:
        out.append("<h2>References</h2>")
        for link in model["links"]:
            out.append(f"<p><a href=\"{escape(link['url'], quote=True)}\">{escape(link['label'])}</a>: {escape(link['url'])}</p>")
            if print_qr:
                data = base64.b64encode(_qr_png(link["url"])).decode("ascii")
                out.append(f"<img class='qr' alt=\"QR code for {escape(link['label'], quote=True)}\" src='data:image/png;base64,{data}'>")
    if _checks(model):
        out += ["<h2>Checks and review</h2>", *[f"<p>{escape(line)}</p>" for line in _checks(model)]]
    return ("\n".join(out) + "</body></html>").encode("utf-8")


def _pdf_fonts():
    import os
    import reportlab
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    directory = os.path.join(os.path.dirname(reportlab.__file__), "fonts")
    for name, filename in (("DanceSans", "Vera.ttf"), ("DanceSansBold", "VeraBd.ttf")):
        if name not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(name, os.path.join(directory, filename)))
    return "DanceSans", "DanceSansBold"


def render_pdf(model, paper="letter", large_print=False, print_qr=False):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, letter
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, LongTable, TableStyle, Image
    regular, bold = _pdf_fonts()
    output = BytesIO()
    size = 16 if large_print else 10.5
    body = ParagraphStyle("Body", fontName=regular, fontSize=size, leading=size * 1.4, spaceAfter=6)
    heading = ParagraphStyle("Heading", parent=body, fontName=bold, fontSize=size + 2, leading=size * 1.5, spaceBefore=12, spaceAfter=7, keepWithNext=True)
    title = ParagraphStyle("Title", parent=heading, fontSize=23 if not large_print else 28, leading=32)
    doc = SimpleDocTemplate(output, pagesize=A4 if paper == "a4" else letter, leftMargin=0.65 * inch, rightMargin=0.65 * inch,
                            topMargin=0.6 * inch, bottomMargin=0.6 * inch, title=model["title"], author=model["choreographer"])
    paragraph = lambda text, style=body: Paragraph(escape(text).replace("\n", "<br/>"), style)
    story = [paragraph(model["title"], title)] + [paragraph(line) for line in _summary(model)]
    for occurrence in model["occurrences"]:
        story.append(paragraph(_occurrence_title(occurrence), heading))
        rows = [[paragraph("Count"), paragraph("Instructions")]]
        rows += [[paragraph(row["label"] + "\n" + row["duration_counts"] + " counts"), paragraph(_row_text(row))] for row in occurrence["rows"]]
        table = LongTable(rows, colWidths=[1.0 * inch, doc.width - inch], repeatRows=1, splitByRow=1, splitInRow=1, hAlign="LEFT")
        table.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eeeeee")),
                                   ("LINEBELOW", (0, 0), (-1, -1), 0.3, colors.HexColor("#cccccc")), ("LEFTPADDING", (0, 0), (-1, -1), 5),
                                   ("RIGHTPADDING", (0, 0), (-1, -1), 5), ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))
        story.extend([table, Spacer(1, 7)])
    if model["links"]:
        story.append(paragraph("References", heading))
        for link in model["links"]:
            text = f'<link href="{escape(link["url"], quote=True)}" color="#174783"><u>{escape(link["label"])}</u></link>: {escape(link["url"])}'
            story.append(Paragraph(text, body))
            if print_qr:
                story.append(Image(BytesIO(_qr_png(link["url"])), width=1.05 * inch, height=1.05 * inch, hAlign="LEFT"))
    if _checks(model):
        story.append(paragraph("Checks and review", heading))
        story += [paragraph(line) for line in _checks(model)]
    def page_footer(canvas, document):
        canvas.setFont(regular, 8)
        canvas.drawRightString(document.pagesize[0] - document.rightMargin, 0.3 * inch, f"Page {document.page}")
    doc.build(story, onFirstPage=page_footer, onLaterPages=page_footer)
    return output.getvalue()


def render_docx(model, paper="letter", large_print=False, print_qr=False):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = (Inches(8.2677), Inches(11.6929)) if paper == "a4" else (Inches(8.5), Inches(11))
    section.top_margin = section.bottom_margin = Inches(0.65)
    section.left_margin = section.right_margin = Inches(0.65)
    for name in ("Normal", "Title", "Heading 1", "Heading 2"):
        style = document.styles[name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor(0, 0, 0)
    document.styles["Normal"].font.size = Pt(16 if large_print else 11)
    document.styles["Heading 2"].font.size = Pt(18 if large_print else 12)
    document.styles["Normal"].paragraph_format.space_after = Pt(6)
    document.add_paragraph(model["title"], style="Title")
    for line in _summary(model):
        document.add_paragraph(line)
    for occurrence in model["occurrences"]:
        document.add_paragraph(_occurrence_title(occurrence), style="Heading 2")
        table = document.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        table.autofit = False
        count_width = Inches(1.4 if large_print else 1.05)
        widths = [count_width, section.page_width - section.left_margin - section.right_margin - count_width]
        for column, width in zip(table.columns, widths):
            column.width = width
        table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Count", "Instructions"
        trpr = table.rows[0]._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:tblHeader"))
        for row in occurrence["rows"]:
            cells = table.add_row().cells
            cells[0].text = row["label"] + "\n" + row["duration_counts"] + " counts"
            cells[1].text = _row_text(row)
        # Word honors cell preferred widths as well as the table grid. Leaving
        # the original equal header widths can push a fixed table off the page.
        for table_row in table.rows:
            table_row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            for cell, width in zip(table_row.cells, widths):
                cell.width = width
        for cell in table.rows[0].cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = True
        document.add_paragraph()
    if model["links"]:
        document.add_paragraph("References", style="Heading 2")
        for link in model["links"]:
            p = document.add_paragraph()
            hyperlink = OxmlElement("w:hyperlink")
            hyperlink.set(qn("r:id"), p.part.relate_to(link["url"], RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
            run, props, color, underline, text = [OxmlElement(tag) for tag in ("w:r", "w:rPr", "w:color", "w:u", "w:t")]
            color.set(qn("w:val"), "174783")
            underline.set(qn("w:val"), "single")
            props.extend([color, underline])
            text.text = link["label"]
            run.extend([props, text]); hyperlink.append(run); p._p.append(hyperlink)
            p.add_run(": " + link["url"])
            if print_qr:
                document.add_picture(BytesIO(_qr_png(link["url"])), width=Inches(1.05))
    if _checks(model):
        document.add_paragraph("Checks and review", style="Heading 2")
        for line in _checks(model):
            document.add_paragraph(line)
    document.core_properties.title = model["title"]
    document.core_properties.author = model["choreographer"]
    output = BytesIO(); document.save(output)
    return output.getvalue()


TABLE_FIELDS = ["title", "choreographer", "song", "artist", "part", "occurrence", "reason", "event_id", "start_count", "end_count", "duration_counts", "count_label", "instructions", "facing_degrees", "support", "status"]


def _table_rows(model):
    yield TABLE_FIELDS
    for occurrence in model["occurrences"]:
        for row in occurrence["rows"]:
            yield [model["title"], model["choreographer"], model["song_title"], model["artist"], occurrence["part_name"],
                   occurrence["occurrence_number"], occurrence.get("reason") or "", row["id"], row["start_count"], row["end_count"],
                   row["duration_counts"], row["label"], _row_text(row), row["facing"] or "unknown", row["support"], model["status"]]


def _spreadsheet_text(value):
    value = str(value)
    return "'" + value if value.lstrip().startswith(("=", "+", "-", "@", "\t", "\r")) else value


def render_csv(model):
    output = StringIO(newline="")
    writer = csv.writer(output)
    writer.writerows([_spreadsheet_text(value) for value in row] for row in _table_rows(model))
    return output.getvalue().encode("utf-8-sig")


def render_xlsx(model):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    workbook = Workbook(); sheet = workbook.active; sheet.title = "Choreography"
    for row in _table_rows(model):
        sheet.append([str(value) for value in row])
        for cell in sheet[sheet.max_row]:
            cell.data_type = "s"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for cell in sheet[1]:
        cell.font = Font(bold=True); cell.fill = PatternFill("solid", fgColor="EEEEEE")
    sheet.freeze_panes = "A2"; sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = 19
    sheet.column_dimensions["M"].width = 65
    links = workbook.create_sheet("References"); links.append(["Label", "URL"])
    for item in model["links"]:
        links.append([item["label"], item["url"]])
        for cell in links[links.max_row]:
            cell.data_type = "s"
        links.cell(links.max_row, 2).hyperlink = item["url"]
    links.column_dimensions["A"].width = 30; links.column_dimensions["B"].width = 80
    output = BytesIO(); workbook.save(output)
    return output.getvalue()


def _milliseconds(value):
    if value < 0:
        raise ExportError("A cue falls before the recording begins. Correct first-count timing or the pickup.")
    return (value.numerator * 2000 + value.denominator) // (2 * value.denominator)


def _stamp(milliseconds, separator):
    hours, rest = divmod(milliseconds, 3600000)
    minutes, rest = divmod(rest, 60000)
    seconds, ms = divmod(rest, 1000)
    return f"{hours:02}:{minutes:02}:{seconds:02}{separator}{ms:03}"


def render_captions(model, format):
    if not model["compiled"]["valid"]:
        raise ExportError("Correct choreography errors before exporting timed cues.")
    mapping = model["music_map"]
    try:
        if "bpm" not in mapping and not mapping.get("anchors"):
            raise ValueError("Set BPM or reviewed timing anchors before exporting captions.")
        if (mapping.get("beat_times") or mapping.get("variable_tempo")) and not mapping.get("anchors"):
            raise ValueError("Add reviewed count/time anchors for variable-tempo captions.")
        mapping = validate_map(mapping)
    except (KeyError, TypeError, ValueError, ZeroDivisionError) as exc:
        raise ExportError(f"Review the music timing before exporting captions: {exc}") from exc
    lines = ["WEBVTT", ""] if format == "vtt" else []
    index, previous_end = 0, -1
    for occurrence in model["occurrences"]:
        for row in occurrence["rows"]:
            # Each boundary comes from the same mapping used by practice. There
            # is no accumulation of rounded subtitle durations across events.
            start = _milliseconds(Fraction(str(count_to_seconds(exact_count(row["start_count"]), mapping))))
            end = _milliseconds(Fraction(str(count_to_seconds(exact_count(row["end_count"]), mapping))))
            if start == end:
                if exact_count(row["duration_counts"]) == 0:
                    continue
                raise ExportError("A cue is shorter than caption millisecond precision; combine adjacent events.")
            if start < previous_end:
                raise ExportError("Caption events overlap; review the choreography timing.")
            previous_end = end; index += 1
            sep = "." if format == "vtt" else ","
            caption = f"{occurrence['part_name']} - {row['label']}: {_row_text(row)}"
            if not model["compiled"]["verified"]:
                caption = "DRAFT - UNVERIFIED MECHANICS | " + caption
            caption = " ".join(caption.replace("-->", "→").split())
            lines += [str(index), f"{_stamp(start, sep)} --> {_stamp(end, sep)}", escape(caption, quote=False), ""]
    if not index:
        raise ExportError("No nonzero timed events are available for captions.")
    return ("\n".join(lines) + "\n").encode("utf-8")


def _media_references(project):
    draft = project.get("draft") or {}
    values = draft.get("attachments", project.get("attachments", [])) or []
    if not isinstance(values, list):
        values = []
    out = []
    for index, item in enumerate(values[:1000]):
        if not isinstance(item, dict):
            continue
        raw_name = item.get("filename") or item.get("original_filename") or item.get("name") or item.get("path") or item.get("local_path") or "Media file"
        filename = re.split(r"[\\/]", _text(raw_name))[-1]
        reference = {"id": _text(item.get("id") or f"media-{index + 1}"), "filename": filename,
                     "caption": _text(item.get("caption")), "needs_relink": True}
        for key in ("sha256", "mime_type", "mime", "bytes", "move_id", "orientation", "start_seconds", "end_seconds", "range_review"):
            if key in item:
                reference[key] = _redact(item[key])
        if safe_url(item.get("url")):
            reference["url"] = safe_url(item["url"])
        out.append(reference)
    return out


def portable_project(project, include_lyrics=False):
    draft, meta, song = _metadata(project)
    compiled = _compile_project(project)
    if not isinstance(draft.get("choreography"), (dict, list)) and not compiled["events"]:
        raise ExportError("The project has no portable choreography yet.")
    # Preserve invalid authored drafts too; normalization must not silently drop
    # missing part references or unsupported data during a backup.
    if compiled["valid"]:
        choreography = _annotated_document(project, compiled)
    else:
        choreography = deepcopy(draft.get("choreography"))
    music = draft.get("music_map") or {}
    if not isinstance(music, dict):
        raise ExportError("Music timing must be an object.")
    portable_meta = {key: _redact(value) for key, value in meta.items() if key in META_FIELDS}
    portable_meta["links"] = _links(meta)
    for key in ("youtube_url", "sheet_url", "song_url", "source_url"):
        if key in portable_meta:
            portable_meta[key] = safe_url(portable_meta[key]) or ""
    if "spotify_url" in portable_meta:
        portable_meta["spotify_url"] = _public_spotify(portable_meta) or ""
    references = _media_references(project)
    for item in _used_annotations(project, compiled):
        for attachment in item.get("attachments", []):
            if not any(saved.get("id") == attachment.get("id") for saved in references):
                references.append(deepcopy(attachment))
    document = {"kind": PACKAGE_KIND, "schema_version": PACKAGE_SCHEMA,
            "name": _text(project.get("name") or meta.get("dance_title") or "Restored dance", 160),
            "song": {"title": _text(song.get("title")), "artist": _text(song.get("artist"))},
            "draft": {"choreography": _redact(choreography), "sheet_meta": portable_meta,
                      "music_map": {key: _redact(value) for key, value in music.items() if key in MUSIC_FIELDS}, "attachments": references},
            "movement_snapshots": _redact(draft.get("movement_snapshots") or project.get("movement_snapshots") or {}),
            "media_references": references,
            "includes_lyrics": include_lyrics is True,
            "omitted": ["Audio and media bytes", "Local file paths", "Provider settings and credentials", "Project revision history", "Analysis caches"],
            "source_revision": project.get("document_revision", 0)}
    if include_lyrics is True:
        lyrics = draft.get("lyrics_raw", project.get("lyrics_raw", ""))
        if not isinstance(lyrics, str):
            raise ExportError("Lyrics must be plain text to include them in a project package.")
        # Preserve exact authored text, including newlines; the package byte
        # limit is checked after serialization. Do not silently truncate it.
        document["draft"]["lyrics_raw"] = lyrics
    else:
        document["omitted"].append("Lyrics")
    return document


def render_package(project, format="zip", include_lyrics=False):
    document = portable_project(project, include_lyrics=include_lyrics)
    payload = _json_bytes(document)
    if len(payload) > MAX_MEMBER_BYTES:
        raise ExportError("Project is too large for a portable package.")
    if format == "json":
        return payload
    manifest = {"kind": PACKAGE_KIND, "schema_version": PACKAGE_SCHEMA,
                "entries": [{"path": "project.json", "size": len(payload), "sha256": hashlib.sha256(payload).hexdigest()}]}
    output = BytesIO()
    # Stored entries keep even highly repetitive legal projects below the reader's
    # compression-ratio limit and are small enough for this media-free format.
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_STORED) as archive:
        archive.writestr("manifest.json", _json_bytes(manifest))
        archive.writestr("project.json", payload)
    return output.getvalue()


def _decode_json(payload):
    def pairs(items):
        result = {}
        for key, value in items:
            if key in result:
                raise ExportError("Duplicate JSON object keys are not allowed.")
            result[key] = value
        return result
    def invalid(value):
        raise ExportError("Nonfinite JSON numbers are not allowed.")
    try:
        return json.loads(payload.decode("utf-8-sig"), object_pairs_hook=pairs, parse_constant=invalid)
    except (UnicodeError, json.JSONDecodeError, RecursionError) as exc:
        raise ExportError("The package contains invalid JSON.") from exc


def _validate_portable(document):
    if not isinstance(document, dict) or document.get("kind") != PACKAGE_KIND:
        raise ExportError("This file is not a Line Dance Creator project package.")
    if type(document.get("schema_version")) is not int or document["schema_version"] != PACKAGE_SCHEMA:
        raise ExportError("This project package schema is unsupported.")
    if not isinstance(document.get("name"), str) or not document["name"].strip() or len(document["name"]) > 160:
        raise ExportError("The package needs a valid project name.")
    draft = document.get("draft")
    if not isinstance(draft, dict) or not isinstance(draft.get("choreography"), (dict, list)):
        raise ExportError("The package is missing its editable choreography.")
    if isinstance(draft["choreography"], dict):
        version = draft["choreography"].get("schema_version", 1)
        if type(version) is not int or version != 1:
            raise ExportError("The choreography schema is unsupported.")
    for key in ("sheet_meta", "music_map"):
        if not isinstance(draft.get(key, {}), dict):
            raise ExportError(f"Package {key} must be an object.")
    if not isinstance(document.get("song", {}), dict) or not isinstance(document.get("movement_snapshots", {}), dict):
        raise ExportError("Package metadata or movement snapshots are malformed.")
    # Project again through the explicit export projection. Unknown fields,
    # paths, connection settings, and caller-selected identity never reach store.
    clean = portable_project({"name": document["name"], "song": document.get("song", {}),
                              "draft": {**draft, "movement_snapshots": document.get("movement_snapshots", {})}},
                             include_lyrics="lyrics_raw" in draft)
    return clean


def read_project_package(payload, filename="project.zip"):
    if not isinstance(payload, bytes) or not payload or len(payload) > MAX_INPUT_BYTES:
        raise ExportError("Choose a nonempty project package no larger than 16 MB.")
    if payload.startswith(b"PK"):
        try:
            with zipfile.ZipFile(BytesIO(payload)) as archive:
                infos = archive.infolist()
                if len(infos) > MAX_MEMBERS:
                    raise ExportError("Too many package members.")
                names, total = set(), 0
                for info in infos:
                    name = info.filename
                    path = PurePosixPath(name)
                    if name in names or name.casefold() in {existing.casefold() for existing in names}:
                        raise ExportError("Duplicate package members are not allowed.")
                    names.add(name)
                    if path.is_absolute() or ".." in path.parts or "\\" in name or ":" in name or name.startswith("/") or name not in {"manifest.json", "project.json"}:
                        raise ExportError("Unexpected or unsafe package path.")
                    if info.is_dir() or stat.S_ISLNK(info.external_attr >> 16) or info.flag_bits & 1:
                        raise ExportError("Directories, symlinks, and encrypted entries are not supported.")
                    total += info.file_size
                    if info.file_size > MAX_MEMBER_BYTES or total > MAX_UNCOMPRESSED_BYTES or info.file_size / max(1, info.compress_size) > MAX_COMPRESSION_RATIO:
                        raise ExportError("Package expansion limits exceeded.")
                if names != {"manifest.json", "project.json"}:
                    raise ExportError("The package needs its manifest and project document.")
                manifest = _decode_json(archive.read("manifest.json"))
                project_bytes = archive.read("project.json")
                if not isinstance(manifest, dict) or manifest.get("kind") != PACKAGE_KIND or type(manifest.get("schema_version")) is not int or manifest.get("schema_version") != PACKAGE_SCHEMA:
                    raise ExportError("The package manifest schema is unsupported.")
                expected = [{"path": "project.json", "size": len(project_bytes), "sha256": hashlib.sha256(project_bytes).hexdigest()}]
                if manifest.get("entries") != expected:
                    raise ExportError("Package integrity check failed.")
                document = _decode_json(project_bytes)
        except (zipfile.BadZipFile, RuntimeError, NotImplementedError, OSError) as exc:
            raise ExportError("The project archive is corrupt or unsupported.") from exc
    else:
        document = _decode_json(payload)
    return _validate_portable(document)


def preview_project_package(payload, filename="project.zip"):
    document = read_project_package(payload, filename)
    result = compile_choreography(document["draft"]["choreography"], document.get("movement_snapshots"))
    return {"name": document["name"], "counts": result["total_counts"], "status": result["status"],
            "media_to_relink": len(document["media_references"]), "omitted": document["omitted"],
            "includes_lyrics": document["includes_lyrics"],
            "lyrics_characters": len(document["draft"].get("lyrics_raw", "")),
            "restores_as": "new_project", "issues": result["issues"]}


def restore_project_package(payload, filename="project.zip"):
    from . import project as store
    document = read_project_package(payload, filename)
    song = document["song"]
    created = store.create_project(document["name"], song.get("title"), song.get("artist"))
    try:
        draft = document["draft"]
        draft["song"] = {"title": song.get("title", ""), "artist": song.get("artist", "")}
        draft["movement_snapshots"] = document.get("movement_snapshots", {})
        store.save_workspace(created["id"], draft, created["document_revision"])
    except Exception:
        # The only removable directory is the newly allocated import project.
        store.delete_project(created["id"])
        raise
    return {"id": created["id"], "name": created["name"], "media_to_relink": len(document["media_references"])}


def export_project(project, format, paper="letter", large_print=False, print_qr=False, include_lyrics=False):
    format = str(format).lower()
    if format not in FORMATS:
        raise ExportError("Unsupported export format.")
    if paper not in ("letter", "a4"):
        raise ExportError("Paper must be letter or a4.")
    if format in ('srt','vtt'):
        from .project_media import recording_review_needed
        if recording_review_needed(project):
            raise ExportError('The recording or timing changed. Review this recording in Music before exporting timed captions.')
    if format in ("zip", "json"):
        output = render_package(project, format, include_lyrics=include_lyrics)
    else:
        model = build_sheet_model(project)
        print_qr = bool(print_qr or (project.get("draft", {}).get("sheet_meta") or {}).get("print_qr"))
        if format in ("pdf", "docx", "html"):
            output = {"pdf": render_pdf, "docx": render_docx, "html": render_html}[format](model, paper, large_print, print_qr)
        elif format in ("srt", "vtt"):
            output = render_captions(model, format)
        else:
            output = {"txt": render_txt, "csv": render_csv, "xlsx": render_xlsx}[format](model)
    mime = {"pdf": "application/pdf", "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "txt": "text/plain; charset=utf-8", "html": "text/html; charset=utf-8", "csv": "text/csv; charset=utf-8", "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "srt": "application/x-subrip", "vtt": "text/vtt; charset=utf-8", "json": "application/json", "zip": "application/zip"}[format]
    title = ((project.get("draft") or {}).get("sheet_meta") or {}).get("dance_title") or project.get("name") or "dance"
    filename = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "-", _text(title, 160)).strip(". ") or "dance"
    return output, mime, filename + "." + format
