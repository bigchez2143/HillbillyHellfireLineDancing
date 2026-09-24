"""S04 local library evidence; fixtures isolate every write from user data."""
import copy
from concurrent.futures import ThreadPoolExecutor
import hashlib
import io
import json
from pathlib import Path
import sqlite3
import sys
import threading
import zipfile

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from PIL import Image

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from engine import library_store as store
from engine.choreography import compile_choreography
from library_api import router


@pytest.fixture(autouse=True)
def library(tmp_path, monkeypatch):
    root = tmp_path / 'library'
    monkeypatch.setattr(store, 'LIBRARY_DIR', root)
    return root


@pytest.fixture
def client():
    app = FastAPI()
    app.include_router(router)
    with TestClient(app) as value:
        yield value


def fields(name='My swing', mechanics=False):
    result = {'name': name, 'aliases': ['My alternate'], 'explanation': 'Step, then hold.',
              'duration_counts': '3/2', 'difficulty': 'Improver'}
    if mechanics:
        result['events'] = [{'duration_counts': '1', 'text': 'Step right', 'support_before': 'L',
                             'support_after': 'R', 'rotation_deg': '0'},
                            {'duration_counts': '1/2', 'text': 'Hold', 'support_before': 'R',
                             'support_after': 'same', 'rotation_deg': '0'}]
    return result


def preview(rows, **kwargs):
    return store.preview_import('moves.json', json.dumps(rows).encode(), **kwargs)


def commit(value, decisions=None):
    return store.commit_import(value['token'], value['digest'], decisions or [{'row': 1, 'action': 'create'}], True)


def png():
    stream = io.BytesIO()
    Image.new('RGB', (5, 5), 'red').save(stream, format='PNG')
    return stream.getvalue()


def test_catalog_preserved_and_custom_storage_is_separate(library):
    catalog = Path(store.__file__).resolve().parents[2] / 'data' / 'step-database.json'
    before = hashlib.sha256(catalog.read_bytes()).hexdigest()
    result = store.get_library()
    assert result['coverage'] == {'reference_records': 216, 'core_patterns': 39, 'custom_records': 0}
    store.create_record(fields())
    assert hashlib.sha256(catalog.read_bytes()).hexdigest() == before
    assert store.get_library()['coverage']['custom_records'] == 1


def test_environment_directory_is_resolved_without_writing_default(tmp_path, monkeypatch):
    monkeypatch.setattr(store, 'LIBRARY_DIR', None)
    monkeypatch.setenv('LINE_DANCE_LIBRARY_DIR', str(tmp_path / 'portable'))
    store.create_record(fields())
    assert (tmp_path / 'portable' / 'library.sqlite3').is_file()


def test_full_record_roundtrip_and_original_snapshot_survives_delete():
    incoming = fields(mechanics=True)
    incoming.update(origins=[{'kind': 'self', 'label': 'My own variation', 'rights_note': 'My instruction text'}],
                    links=[{'url': 'https://example.org/lesson', 'label': 'Lesson'}])
    record = store.create_record(incoming)
    frozen = store.snapshot(record['id'], 1)
    assert frozen['duration_counts'] == '3/2'
    assert store.get_record(record['id'])['events'] == incoming['events']
    store.update_record(record['id'], 1, {'explanation': 'Changed'})
    store.delete_record(record['id'], 2)
    assert store.snapshot(record['id'], 1) == frozen
    assert not store.get_library()['records']
    assert store.get_library(True)['records'][0]['deleted']
    with pytest.raises(store.Conflict):
        store.update_record(record['id'], 3, {'name': 'Cannot edit deleted'})


def test_unknown_mechanics_remain_unverified_but_portable():
    record = store.create_record(fields())
    frozen = store.snapshot(record['id'])
    report = compile_choreography([frozen])
    assert report['status'] == 'UNVERIFIED'
    assert not any(i['code'] == 'SNAPSHOT_MISSING' for i in report['issues'])
    assert record['review']['status'] == 'UNVERIFIED'
    with pytest.raises(ValueError, match='mechanics'):
        store.review_record(record['id'], 1, 'Me', '', True)


def test_author_review_requires_confirmation_and_complete_mechanics_then_edit_resets():
    record = store.create_record(fields(mechanics=True))
    with pytest.raises(ValueError, match='confirmation'):
        store.review_record(record['id'], 1, 'Me', '', False)
    reviewed = store.review_record(record['id'], 1, 'Me', 'Checked foot changes', True)
    assert reviewed['review']['status'] == 'AUTHOR_REVIEWED'
    assert reviewed['generator_eligible'] is False
    changed = store.update_record(record['id'], 2, {'duration_counts': '2'})
    assert changed['review'] == {'status': 'UNVERIFIED'}
    assert store.snapshot(record['id'], 2)['review']['reviewer'] == 'Me'


@pytest.mark.parametrize('bad', [True, '-1', 'nan', 'inf', '1/0', {'numerator': 1, 'denominator': 0}, '100001'])
def test_invalid_exact_counts_are_rejected(bad):
    with pytest.raises(ValueError):
        store.create_record({**fields(), 'duration_counts': bad})


def test_missing_count_and_ambiguous_mechanics_preserved_without_invention():
    record = store.create_record({'name': 'A new styling idea', 'explanation': 'Small shoulder motion.'})
    assert record['duration_counts'] is None
    assert record['events'] == []
    assert 'duration_counts' not in store.snapshot(record['id'])


def test_two_concurrent_editors_only_one_can_commit():
    record = store.create_record(fields())
    barrier = threading.Barrier(2)
    def edit(name):
        barrier.wait()
        try:
            return store.update_record(record['id'], 1, {'name': name})['name']
        except store.Conflict:
            return 'conflict'
    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(edit, ['One', 'Two']))
    assert results.count('conflict') == 1
    assert store.get_record(record['id'])['version'] == 2
    assert store.get_record(record['id'], 1)['name'] == 'My swing'


def test_preview_never_mutates_records_and_reports_alias_and_internal_duplicates():
    existing = store.create_record(fields('Existing'))
    result = preview([fields('Fresh'), fields('Fresh')])
    assert result['rows'][0]['duplicates'] == [{'kind': 'user', 'id': existing['id'], 'name': 'Existing', 'version': 1}]
    assert result['rows'][1]['duplicate_rows'] == [1]
    assert len(store.get_library()['records']) == 1


def test_duplicate_core_or_reference_requires_explicit_variant():
    result = preview([{'name': 'Grapevine', 'explanation': 'My variant.'}])
    assert result['rows'][0]['duplicates']
    with pytest.raises(store.Conflict, match='variant'):
        commit(result)
    applied = commit(result, [{'row': 1, 'action': 'variant'}])
    assert applied['records'][0]['review']['status'] == 'UNVERIFIED'


def test_ai_claims_discarded_and_explicit_user_approval_required():
    result = preview([{**fields(), 'review': {'status': 'VERIFIED'}, 'generator_eligible': True}], origin_kind='ai_user_supplied')
    assert result['rows'][0]['warnings']
    with pytest.raises(ValueError, match='approval'):
        store.commit_import(result['token'], result['digest'], [{'row': 1, 'action': 'create'}], False)
    applied = commit(result)
    record = applied['records'][0]
    assert record['review']['status'] == 'UNVERIFIED'
    assert not record['generator_eligible']
    assert record['origins'][-1]['kind'] == 'ai_user_supplied'


def test_import_digest_and_retry_decisions_cannot_change():
    value = preview([fields()])
    with pytest.raises(store.Conflict, match='preview changed'):
        store.commit_import(value['token'], 'wrong', [{'row': 1, 'action': 'create'}], True)
    first = commit(value)
    assert commit(value) == first
    with pytest.raises(store.Conflict, match='different decisions'):
        commit(value, [{'row': 1, 'action': 'skip'}])
    assert len(store.get_library()['records']) == 1


def test_import_batch_rejects_second_conflict_without_partial_writes():
    value = preview([{'name': 'Unique local A'}, {'name': 'Unique local A'}])
    with pytest.raises(store.Conflict):
        commit(value, [{'row': 1, 'action': 'create'}, {'row': 2, 'action': 'create'}])
    assert not store.get_library()['records']
    result = commit(value, [{'row': 1, 'action': 'create'}, {'row': 2, 'action': 'variant'}])
    assert result['applied'] == 2


def test_import_update_cas_and_rollback_keep_all_saved_versions():
    original = store.create_record(fields())
    value = preview([{**fields(), 'explanation': 'Imported edit'}])
    with pytest.raises(store.Conflict):
        commit(value, [{'row': 1, 'action': 'update', 'target_id': original['id'], 'expected_version': 9}])
    result = commit(value, [{'row': 1, 'action': 'update', 'target_id': original['id'], 'expected_version': 1}])
    assert store.get_record(original['id'])['explanation'] == 'Imported edit'
    store.rollback_batch(result['batch_id'], True)
    assert store.get_record(original['id'])['explanation'] == original['explanation']
    assert store.get_record(original['id'])['version'] == 3
    assert store.get_record(original['id'], 2)['explanation'] == 'Imported edit'
    assert store.rollback_batch(result['batch_id'], True)['state'] == 'rolled_back'


def test_rollback_refuses_all_if_any_imported_record_changed_later():
    result = commit(preview([{'name': 'One'}, {'name': 'Two'}]), [{'row': 1, 'action': 'create'}, {'row': 2, 'action': 'create'}])
    one, two = result['records']
    store.update_record(two['id'], 1, {'explanation': 'Keep my work'})
    with pytest.raises(store.Conflict):
        store.rollback_batch(result['batch_id'], True)
    assert not store.get_record(one['id'])['deleted']
    assert store.get_record(two['id'])['explanation'] == 'Keep my work'


@pytest.mark.parametrize('format_name', ['json', 'csv', 'xlsx'])
def test_templates_are_readable_roundtrips(format_name):
    payload, mime = store.template(format_name)
    result = store.preview_import('template.' + format_name, payload)
    assert not result['rows'][0]['errors']
    assert result['rows'][0]['fields']['duration_counts'] == '3/2'
    assert mime


def test_csv_mapping_multiline_and_exact_count():
    payload = b'title,description,beats\r\n"New pattern","Step right,\nthen left",1.5\r\n'
    result = store.preview_import('mapped.csv', payload, {'title': 'name', 'description': 'explanation', 'beats': 'duration_counts'})
    fields = result['rows'][0]['fields']
    assert fields['explanation'] == 'Step right,\nthen left'
    assert fields['duration_counts'] == '3/2'


@pytest.mark.parametrize('filename,payload', [('bad.json', b'{'), ('bad.csv', b'name,name\nA,B'),
    ('bad.xlsx', b'fake archive'), ('bad.exe', b'no'), ('../bad.json', b'[]'), ('bad.json', b'[]'),
    ('bad.json', b'[{"name":"X","duration_counts":NaN}]')])
def test_malformed_imports_rejected(filename, payload):
    with pytest.raises(ValueError):
        store.preview_import(filename, payload)


def test_spreadsheet_formulas_and_archive_traversal_rejected():
    from openpyxl import Workbook
    workbook = Workbook()
    workbook.active.append(['name', 'duration_counts'])
    workbook.active.append(['Formula move', '=1+1'])
    stream = io.BytesIO()
    workbook.save(stream)
    with pytest.raises(ValueError, match='formulas'):
        store.preview_import('formula.xlsx', stream.getvalue())
    archive = io.BytesIO()
    with zipfile.ZipFile(archive, 'w') as writer:
        writer.writestr('../outside', 'unsafe')
    with pytest.raises(ValueError, match='archive entry'):
        store.preview_import('traversal.xlsx', archive.getvalue())


def test_secret_fields_never_persist_even_when_unmapped(library):
    with pytest.raises(ValueError, match='credentials'):
        preview([{'name': 'Proposed', 'api_key': 'do-not-store-this'}])
    assert not library.exists()
    for url in ['https://user:secret@example.org/file', 'https://example.org/?access_token=secret', 'file:///C:/secret']:
        with pytest.raises(ValueError):
            store.create_record({**fields(), 'links': [url]})


def test_bad_rows_are_visible_and_cannot_be_approved():
    value = preview([{'name': 'Impossible count', 'duration_counts': '-4'}, {'name': 'Okay'}])
    assert value['rows'][0]['errors']
    with pytest.raises(ValueError, match='Row 1 has errors'):
        commit(value)
    applied = commit(value, [{'row': 1, 'action': 'skip'}, {'row': 2, 'action': 'create'}])
    assert applied['applied'] == 1


def test_attachment_original_and_frozen_metadata_survive_soft_delete():
    record = store.create_record(fields())
    original = png()
    attached = store.attach_media(record['id'], 1, 'demonstration.png', original, 'Right foot', 'front')
    meta = attached['attachment']
    path, stored = store.media_content(meta['id'])
    assert path.read_bytes() == original
    assert stored['sha256'] == hashlib.sha256(original).hexdigest()
    assert attached['record']['review']['status'] == 'UNVERIFIED'
    assert store.snapshot(record['id'], 1)['attachments'] == []
    store.delete_record(record['id'], 2)
    assert store.snapshot(record['id'], 2)['attachments'][0] == meta
    assert store.media_content(meta['id'])[0].is_file()


@pytest.mark.parametrize('name,payload', [('../x.png', b'a'), ('C:\\x.png', b'a'),
                                        ('x.svg', b'<svg/>'), ('x.png', b'<html>'), ('x.mp4', b'not video')])
def test_unsafe_or_mislabeled_media_rejected(name, payload):
    record = store.create_record(fields())
    with pytest.raises(ValueError):
        store.attach_media(record['id'], 1, name, payload)
    assert store.get_record(record['id'])['version'] == 1


def test_attachment_conflict_or_storage_failure_leaves_no_new_original(library, monkeypatch):
    record = store.create_record(fields())
    store.update_record(record['id'], 1, {'explanation': 'Later'})
    with pytest.raises(store.Conflict):
        store.attach_media(record['id'], 1, 'x.png', png())
    assert not list((library / 'media').glob('*'))
    def failed(*args):
        raise OSError('Disk write failed')
    monkeypatch.setattr(store, '_put', failed)
    with pytest.raises(OSError):
        store.attach_media(record['id'], 2, 'x.png', png())
    assert not list((library / 'media').glob('*'))
    assert store.get_record(record['id'])['version'] == 2


def test_media_path_tampering_cannot_serve_file_outside_library(library):
    record = store.create_record(fields())
    meta = store.attach_media(record['id'], 1, 'x.png', png())['attachment']
    outside = library.parent / 'private.txt'
    outside.write_text('private')
    with sqlite3.connect(library / 'library.sqlite3') as conn:
        conn.execute('UPDATE media SET filename=? WHERE id=?', ('../../private.txt', meta['id']))
    with pytest.raises(FileNotFoundError):
        store.media_content(meta['id'])


def test_api_contract_rejects_stale_and_bool_versions(client):
    result = client.post('/api/creator/library/records', json={'fields': fields()})
    assert result.status_code == 201
    record = result.json()
    url = '/api/creator/library/records/' + record['id']
    assert client.patch(url, json={'expected_version': True, 'fields': {'name': 'Bad'}}).status_code == 422
    assert client.patch(url, json={'expected_version': 1, 'fields': {'name': 'Good'}}).status_code == 200
    assert client.patch(url, json={'expected_version': 1, 'fields': {'name': 'Lost'}}).status_code == 409
    assert client.get(url + '/snapshot?version=1').json()['name'] == 'My swing'


def test_api_upload_and_template_and_snapshot(client):
    result = client.post('/api/creator/library/imports/preview', files={'file': ('test.json', b'[{"name":"Test uploaded"}]', 'application/json')})
    assert result.status_code == 200, result.text
    value = result.json()
    result = client.post('/api/creator/library/imports/' + value['token'] + '/commit',
                         json={'digest': value['digest'], 'decisions': [{'row': 1, 'action': 'create'}], 'confirmed': True})
    assert result.status_code == 200, result.text
    record = result.json()['records'][0]
    result = client.post('/api/creator/library/records/' + record['id'] + '/attachments',
                         data={'expected_version': '1', 'caption': 'Sample'}, files={'file': ('safe.png', png(), 'image/png')})
    assert result.status_code == 201, result.text
    media_id = result.json()['attachment']['id']
    content = client.get('/api/creator/library/attachments/' + media_id + '/content')
    assert content.content == png()
    assert content.headers['x-content-type-options'] == 'nosniff'
    assert client.get('/api/creator/library/templates/xlsx').status_code == 200


def test_data_root_and_dedicated_library_override(tmp_path, monkeypatch):
    monkeypatch.setattr(store, 'LIBRARY_DIR', None)
    monkeypatch.delenv('LINE_DANCE_LIBRARY_DIR', raising=False)
    monkeypatch.setenv('LINE_DANCE_DATA_DIR', str(tmp_path / 'data'))
    assert store.directory() == tmp_path / 'data' / 'library'
    monkeypatch.setenv('LINE_DANCE_LIBRARY_DIR', str(tmp_path / 'separate'))
    assert store.directory() == tmp_path / 'separate'


def test_import_failure_rolls_back_database_transaction(library, monkeypatch):
    value = preview([{'name': 'Transaction one'}, {'name': 'Transaction two'}])
    real_put = store._put
    written = []
    def fail_second(conn, record):
        written.append(record['id'])
        if len(written) == 2:
            raise OSError('Simulated disk failure')
        return real_put(conn, record)
    monkeypatch.setattr(store, '_put', fail_second)
    with pytest.raises(OSError):
        commit(value, [{'row': 1, 'action': 'create'}, {'row': 2, 'action': 'create'}])
    assert not store.get_library()['records']
    with sqlite3.connect(library / 'library.sqlite3') as conn:
        assert conn.execute('SELECT count(*) FROM versions').fetchone()[0] == 0
        assert conn.execute('SELECT result FROM previews').fetchone()[0] is None
    monkeypatch.setattr(store, '_put', real_put)
    assert commit(value, [{'row': 1, 'action': 'create'}, {'row': 2, 'action': 'create'}])['applied'] == 2


def test_spreadsheet_dimensions_and_zip_bomb_guard(monkeypatch):
    from openpyxl import Workbook
    workbook = Workbook()
    workbook.active.append(['name'])
    workbook.active.cell(1002, 1, 'Far away')
    stream = io.BytesIO()
    workbook.save(stream)
    with pytest.raises(ValueError, match='1000 data rows'):
        store.preview_import('dimensions.xlsx', stream.getvalue())
    monkeypatch.setattr(store, 'MAX_ARCHIVE_BYTES', 10)
    with pytest.raises(ValueError, match='archive is too large'):
        store.preview_import('oversized.xlsx', stream.getvalue())


def test_mapping_collision_does_not_silently_choose_one_column():
    value = store.preview_import('columns.csv', b'name,title\nOriginal,Override\n', {'title': 'name'})
    assert 'same field' in value['rows'][0]['errors'][0]


def test_upload_request_limits_with_and_without_content_length(client, monkeypatch):
    monkeypatch.setattr(store, 'MAX_IMPORT_BYTES', 20)
    oversized = b'x' * 70000
    endpoint = '/api/creator/library/imports/preview'
    assert client.post(endpoint, content=oversized, headers={'Content-Type': 'application/octet-stream'}).status_code == 413
    # Streaming clients need the same bound even when no Content-Length exists.
    def stream():
        yield oversized[:40000]
        yield oversized[40000:]
    response = client.post('/api/creator/library/records', content=stream(), headers={'Content-Type': 'application/json'})
    assert response.status_code == 413
    assert not store.get_library()['records']


def test_media_limit_and_invalid_time_ranges_preserve_record(monkeypatch):
    record = store.create_record(fields())
    monkeypatch.setattr(store, 'MAX_IMAGE_BYTES', 1)
    with pytest.raises(ValueError, match='limit'):
        store.attach_media(record['id'], 1, 'large.png', png())
    # A supplied range is a reference annotation, not a claim the clip was decoded.
    video = b'\x00\x00\x00\x18ftypmp42' + b'\x00' * 12
    with pytest.raises(ValueError, match='endpoints'):
        store.attach_media(record['id'], 1, 'clip.mp4', video, start_seconds='2', end_seconds='1')
    attached = store.attach_media(record['id'], 1, 'clip.mp4', video, start_seconds='1/2', end_seconds='3/2')
    assert attached['attachment']['range_review'] == 'USER_ENTERED_NOT_MEASURED'
    assert attached['attachment']['start_seconds'] == '1/2'


def test_restore_soft_deleted_move_appends_version_and_preserves_frozen_history():
    record = store.create_record(fields(mechanics=True))
    frozen = store.snapshot(record['id'], 1)
    deleted = store.delete_record(record['id'], 1)
    with pytest.raises(store.Conflict):
        store.restore_record(record['id'], 1)
    restored = store.restore_record(record['id'], 2)
    assert restored['version'] == 3 and not restored['deleted']
    assert store.get_record(record['id'], 2) == deleted
    assert store.snapshot(record['id'], 1) == frozen
    assert store.restore_record(record['id'], 3) == restored


def test_all_catalog_reference_snapshots_preserve_names_facts_and_remain_unverified():
    reference, _ = store.reference_catalog()
    assert len(reference) == 216
    for entry in reference:
        frozen = store.reference_snapshot(entry['id'])
        assert frozen['name'] == entry['name']
        assert frozen['explanation'] == entry['description']
        assert frozen['reference_record'] == entry
        assert frozen['reported_mechanics']['foot_start'] == entry['foot_start']
        assert frozen['review']['status'] == 'UNVERIFIED'
        assert not frozen['generator_eligible']
        if entry['counts'] is not None:
            assert frozen['duration_counts'] == str(entry['counts'])
            assert compile_choreography([frozen])['status'] == 'UNVERIFIED'
        else:
            assert frozen['requires_counts']
            assert 'duration_counts' not in frozen
            assert not compile_choreography([frozen])['verified']


def test_reference_manual_count_override_and_snapshot_immutability():
    initial = store.reference_snapshot('reference-0')
    assert initial['name'] == 'Balance' and initial['requires_counts']
    frozen = store.reference_snapshot('reference-0', '1.5')
    assert frozen['duration_counts'] == '3/2'
    assert frozen['count_origin'] == 'USER_SUPPLIED'
    assert frozen['reference_record']['counts'] is None
    assert frozen['snapshot_id'] != store.reference_snapshot('reference-0', '2')['snapshot_id']
    assert compile_choreography([frozen])['status'] == 'UNVERIFIED'
    frozen['reference_record']['name'] = 'Changed copy'
    assert store.reference_snapshot('reference-0') == initial


@pytest.mark.parametrize('reference_id', ['reference-999', '../0', 'reference--1', 'reference-01', 'reference-0/anything'])
def test_invalid_reference_ids_cannot_read_unrelated_files(reference_id):
    with pytest.raises((ValueError, FileNotFoundError)):
        store.reference_snapshot(reference_id)


def test_reference_and_restore_routes(client):
    frozen = client.get('/api/creator/library/reference/reference-0/snapshot?duration_counts=4')
    assert frozen.status_code == 200 and frozen.json()['duration_counts'] == '4'
    record = client.post('/api/creator/library/records', json={'fields': fields()}).json()
    path = '/api/creator/library/records/' + record['id']
    assert client.request('DELETE', path, json={'expected_version': 1}).status_code == 200
    assert client.post(path + '/restore', json={'expected_version': 2}).json()['deleted'] is False
