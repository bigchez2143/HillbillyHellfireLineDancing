"""Real bundled moves through local endpoints and portable/export projections.

All projects use temporary directories. TestClient runs in process; no browser,
network upload, actual song, remote provider, or original project is involved.
"""
from copy import deepcopy
from html import unescape
from io import BytesIO
import json
from pathlib import Path
import sys
import zipfile

import pytest
from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfReader

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from engine import project as store, steps, move_expansion as expansion
from engine.choreography import compile_choreography
from engine.creator_exports import build_sheet_model, export_project, read_project_package
import server

WIZARD = 'expansion-wizard-right-diagonal'
UNKNOWN = 'expansion-moonwalk-provisional-cue'
FROZEN_FIELDS = ('source_hash', 'definition_hash', 'definition_id', 'move_id', 'snapshot_id',
                 'sources', 'source_ids', 'review_status', 'review', 'generator_eligible',
                 'in_generator', 'manual_only', 'mechanically_complete', 'source_pack',
                 'group', 'family', 'level', 'review_note', 'provenance_note',
                 'required_start_facing', 'start_free_foot', 'net_rotation_deg', 'travel')


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(store, 'PROJECTS_DIR', str(tmp_path / 'projects'))
    monkeypatch.setattr(steps, 'CUSTOM_MOVES_PATH', str(tmp_path / 'custom.json'))
    with TestClient(server.app) as value:
        yield value


def document(mid=WIZARD, lead='R'):
    move = expansion.variant(mid, lead)
    move['id'] = 'authored-occurrence'
    return {'schema_version': 1,
            'start': {'free_foot': move['start_free_foot'] or 'R',
                      'facing_deg': move['required_start_facing'] or '0'},
            'parts': [{'id': 'main', 'name': 'Main part', 'moves': [move]}],
            'routine': [{'id': 'run', 'part_id': 'main'}]}


def saved(client, doc=None):
    doc = document() if doc is None else doc
    project = store.create_project('Synthetic expansion practice')
    draft = {'choreography': doc, 'music_map': {'bpm': 120, 'first_count': 2},
             'sheet_meta': {'dance_title': 'Synthetic expansion practice', 'choreographer': 'Test author'}}
    response = client.put(f"/api/projects/{project['id']}/workspace",
                          json={'draft': draft, 'expected_revision': project['document_revision']})
    assert response.status_code == 200, response.text
    return project['id'], response.json()


def exported(client, pid, format):
    response = client.get(f'/api/creator/projects/{pid}/export/{format}')
    assert response.status_code == 200, response.text
    return response.content


def test_creator_moves_endpoint_adds_manual_pack_without_changing_legacy_steps(client):
    legacy = client.get('/api/steps')
    creator = client.get('/api/creator/moves')
    assert legacy.status_code == creator.status_code == 200
    legacy_rows = legacy.json()
    assert legacy_rows == steps.library_json() and len(legacy_rows) == 78
    assert all('events' not in row and not row['move_id'].startswith('expansion-') for row in legacy_rows)
    rows = creator.json()['moves']
    assert rows[:78] == legacy_rows
    added = rows[78:]
    assert len(added) == 86 and len({row['group'] for row in added}) == 22
    assert all(row['events'] and not row['generator_eligible'] for row in added)
    assert sum(row['mechanically_complete'] for row in added) == 78


@pytest.mark.parametrize('mid,status', [(WIZARD, 'VALID'), (UNKNOWN, 'UNVERIFIED')])
def test_save_reopen_accept_preserves_explicit_graph_and_later_draft_does_not_replace_it(client, mid, status):
    doc = document(mid)
    pid, workspace = saved(client, doc)
    reopened = client.get(f'/api/projects/{pid}/workspace').json()
    assert reopened['draft']['choreography'] == doc
    assert reopened['snapshot_issues'] == {'accepted': [], 'draft': []}
    response = client.post(f'/api/creator/projects/{pid}/accept',
                           json={'expected_revision': workspace['document_revision']})
    assert response.status_code == 200, response.text
    assert response.json()['legacy_compatible'] is False
    assert response.json()['validation'] == status
    accepted = store.load_project(pid)
    assert accepted['dance']['source'] == 'structured'
    assert accepted['accepted_choreography_source'] == doc
    assert accepted['dance']['compiler']['status'] == status
    later = deepcopy(doc)
    later['parts'][0]['moves'][0]['events'][0]['text'] = 'Later synthetic draft wording'
    store.save_workspace(pid, {'choreography': later}, accepted['document_revision'])
    after = store.load_project(pid)
    for key in ('accepted_choreography', 'accepted_choreography_source', 'accepted_choreography_hash', 'dance'):
        assert after[key] == accepted[key]


@pytest.mark.parametrize('format', ['json', 'zip'])
@pytest.mark.parametrize('mid,status', [(WIZARD, 'VALID'), (UNKNOWN, 'UNVERIFIED')])
def test_export_import_keeps_frozen_source_and_events_after_pack_changes(client, tmp_path, monkeypatch, format, mid, status):
    doc = document(mid)
    pid, _ = saved(client, doc)
    original = doc['parts'][0]['moves'][0]
    before = compile_choreography(doc)
    payload = exported(client, pid, format)
    portable = read_project_package(payload)
    carried = portable['draft']['choreography']['parts'][0]['moves'][0]
    for field in FROZEN_FIELDS:
        assert carried.get(field) == original.get(field), field
    assert {link['url'] for link in carried['links']} == {source['url'] for source in original['sources']}
    assert compile_choreography(portable['draft']['choreography'])['events'] == before['events']

    changed = expansion.load_pack()
    changed_row = next(row for row in changed['moves'] if row['id'] == mid)
    changed_row['events'][0]['text'] = 'New bundled wording must not replace saved events'
    changed_path = tmp_path / 'changed-expanded.json'
    changed_path.write_text(json.dumps(changed), encoding='utf-8')
    monkeypatch.setattr(expansion, 'PACK_PATH', changed_path)
    assert expansion.variant(mid)['definition_hash'] != original['definition_hash']
    # Import is a local TestClient fixture, never a browser/network file upload.
    preview = client.post('/api/creator/import/preview', files={'file': ('synthetic.' + format, payload)})
    assert preview.status_code == 200 and preview.json()['status'] == status
    restored = client.post('/api/creator/import', files={'file': ('synthetic.' + format, payload)})
    assert restored.status_code == 200, restored.text
    new_pid = restored.json()['id']
    assert new_pid != pid
    new_project = store.load_project(new_pid)
    imported = new_project['draft']['choreography']['parts'][0]['moves'][0]
    for field in FROZEN_FIELDS:
        assert imported.get(field) == original.get(field), field
    compiled = compile_choreography(new_project['draft']['choreography'])
    assert compiled['events'] == before['events'] and compiled['status'] == status
    assert new_project['dance'] == {}
    assert store.load_project(pid)['draft']['choreography'] == doc


@pytest.mark.parametrize('format', ['json', 'zip'])
def test_source_metadata_projection_keeps_public_fields_and_drops_private_fields(client, format):
    doc = document()
    move = doc['parts'][0]['moves'][0]
    move['api_key'] = 'PRIVATE-KEY-SENTINEL'
    move['path'] = r'C:\Private\move.json'
    move['review']['credentials'] = {'secret': 'PRIVATE-REVIEW-SENTINEL'}
    move['sources'][0].update(api_key='PRIVATE-SOURCE-SENTINEL',
                              path=r'C:\Private\reference.pdf',
                              provider_settings={'secret': 'PRIVATE-PROVIDER-SENTINEL'})
    move['sources'].extend([
        {'id': 'local', 'title': 'Local reference', 'url': 'file:///C:/Private/reference.pdf'},
        {'id': 'credentialed', 'title': 'Private account', 'url': 'https://user:PRIVATE-URL-SENTINEL@example.org/'},
        {'id': 'script', 'title': 'Invalid', 'url': 'javascript:alert(1)'},
    ])
    pid, _ = saved(client, doc)
    payload = exported(client, pid, format)
    portable = read_project_package(payload)
    restored = portable['draft']['choreography']['parts'][0]['moves'][0]
    text = json.dumps(portable)
    assert 'PRIVATE-' not in text and 'C:\\\\Private' not in text
    assert 'file:///' not in text and 'javascript:' not in text
    assert restored['source_hash'] == move['source_hash']
    assert restored['definition_hash'] == move['definition_hash']
    assert restored['source_ids'] == move['source_ids']
    assert all(set(source) <= {'id', 'title', 'url'} for source in restored['sources'])
    assert len(restored['sources']) == 3


def test_pdf_docx_and_srt_keep_real_dorothy_late_and_timing_and_source_links(client):
    pid, _ = saved(client)
    pdf = PdfReader(BytesIO(exported(client, pid, 'pdf')))
    pdf_text = '\n'.join(page.extract_text() for page in pdf.pages)
    assert '2&' in pdf_text and '1/2 counts' in pdf_text
    assert 'instructor review required' in pdf_text
    pdf_links = {str(a.get_object().get('/A', {}).get('/URI')) for page in pdf.pages for a in page.get('/Annots', [])}
    assert {source['url'] for source in expansion.variant(WIZARD)['sources']} <= pdf_links

    payload = exported(client, pid, 'docx')
    docx = Document(BytesIO(payload))
    counts = [row.cells[0].text for row in docx.tables[0].rows[1:]]
    assert counts == ['1\n1 counts', '2\n1/2 counts', '2&\n1/2 counts']
    with zipfile.ZipFile(BytesIO(payload)) as archive:
        rels = archive.read('word/_rels/document.xml.rels').decode()
        assert 'TargetMode="External"' in rels
        assert 'ntadance.com' in rels
    srt = unescape(exported(client, pid, 'srt').decode())
    assert '00:00:02,000 --> 00:00:02,500' in srt
    assert '00:00:02,500 --> 00:00:02,750' in srt
    assert '00:00:02,750 --> 00:00:03,000' in srt
    assert 'Main part - 1:' in srt and 'Main part - 2:' in srt and 'Main part - 2&:' in srt
    assert 'Main part - 1&:' not in srt


@pytest.mark.parametrize('format', ['pdf', 'docx', 'srt', 'vtt'])
def test_provisional_real_model_discloses_unverified_in_sheets_and_timed_cues(client, format):
    pid, _ = saved(client, document(UNKNOWN))
    payload = exported(client, pid, format)
    if format == 'pdf':
        text = '\n'.join(page.extract_text() for page in PdfReader(BytesIO(payload)).pages)
    elif format == 'docx':
        doc = Document(BytesIO(payload))
        text = '\n'.join(p.text for p in doc.paragraphs)
    else:
        text = payload.decode()
    assert 'UNVERIFIED MECHANICS' in text
    assert 'modeled checks pass' not in text


def test_sheet_references_can_use_preserved_sources_without_duplicate_links(client):
    doc = document()
    move = doc['parts'][0]['moves'][0]
    move.pop('links')
    pid, _ = saved(client, doc)
    model = build_sheet_model(store.load_project(pid))
    assert {link['url'] for link in model['links']} == {source['url'] for source in move['sources']}
