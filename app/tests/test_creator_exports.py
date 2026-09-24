"""S07 export/portable-import tests using only synthetic authored material."""
from copy import deepcopy
import csv
import hashlib
from io import BytesIO, StringIO
import json
import os
import stat
import sys
import zipfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import pytest
from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader

from engine import project as store
from engine.creator_exports import (ExportError, build_sheet_model, export_project, portable_project,
                                   read_project_package, restore_project_package, preview_project_package,
                                   safe_url, MAX_COMPRESSION_RATIO)
from export_api import router


def fixture_project():
    def movement(identifier, duration, before, after, rotation="0"):
        return {"id": identifier, "name": identifier, "duration_counts": duration,
                "events": [{"id": "action", "duration_counts": duration, "text": f"{identifier}: step with care & keep the count clear.",
                            "support_before": before, "support_after": after, "rotation_deg": rotation}]}
    return {"id": "existing", "name": "Evening Practice", "document_revision": 7,
            "song": {"title": "Synthetic Practice Track", "artist": "Example Musicians", "path": r"C:\Private\unreleased.wav"},
            "api_key": "DO-NOT-EXPORT-TOP-SECRET",
            "draft": {"choreography": {"schema_version": 1,
                        "parts": [{"id": "A", "name": "Basic part", "moves": [movement("Right step", "1", "L", "R"), movement("Left step", "1", "R", "L")]},
                                  {"id": "T", "name": "Half-count tag", "kind": "tag", "moves": [movement("Hold", "1/2", "L", "L")]},
                                  {"id": "E", "name": "Diagonal finish", "kind": "ending", "moves": [movement("Finish", "1/3", "R", "R", "45")]}],
                        "routine": [{"id": "first", "part_id": "A", "repeat": 2}, {"id": "tag", "part_id": "T"},
                                    {"id": "restart", "part_id": "A", "end_after_counts": "1", "reason": "restart"}, {"id": "ending", "part_id": "E"}]},
                      "music_map": {"bpm": 120, "first_count": 2, "meter": 4, "key": "C major", "api_key": "DO-NOT-EXPORT-MAP-SECRET"},
                      "sheet_meta": {"dance_title": "Evening Practice", "choreographer": "Alex Example", "country": "USA", "level_label": "Beginner",
                                     "contact": "alex@example.org", "description": "Read carefully <and> count together.", "credits": "Movements authored for this test",
                                     "youtube_url": "https://example.org/demo", "sheet_url": "https://example.org/sheet", "api_key": "DO-NOT-EXPORT-META-SECRET"},
                      "connection": {"api_key": "DO-NOT-EXPORT-DRAFT-SECRET"}, "lyrics_raw": "PRIVATE LYRICS OMITTED",
                      "attachments": [{"id": "feet-demo", "path": r"C:\Private\foot-demo.mp4", "caption": "Rear view", "sha256": "a" * 64}]}}


@pytest.mark.parametrize("format", ["txt", "html", "csv", "xlsx", "pdf", "docx", "srt", "vtt", "json", "zip"])
def test_every_requested_format_has_output_from_shared_compiler(format):
    project = fixture_project()
    model = build_sheet_model(project)
    assert model["compiled"]["verified"]
    assert model["compiled"]["total_counts"] == "35/6"
    payload, mime, filename = export_project(project, format)
    assert isinstance(payload, bytes) and len(payload) > 20
    assert filename == "Evening Practice." + format
    assert mime


def test_text_and_html_preserve_structure_credits_annotations_and_escape_markup():
    project = fixture_project()
    text = export_project(project, "txt")[0].decode()
    html = export_project(project, "html", paper="a4", large_print=True)[0].decode()
    for value in ("Alex Example", "Synthetic Practice Track", "Half-count tag", "Diagonal finish", "restart after 1 counts", "35/6", "Movements authored for this test"):
        assert value in text and value in html
    assert "&lt;and&gt;" in html
    assert 'href="https://example.org/demo"' in html
    assert "size:a4" in html and "16pt" in html
    assert "Discover Hillbilly Hellfire" not in text + html
    assert "DO-NOT-EXPORT" not in text + html


def test_pdf_text_paper_and_external_link_annotations_are_real():
    pdf = PdfReader(BytesIO(export_project(fixture_project(), "pdf", paper="a4", print_qr=True)[0]))
    text = "\n".join(page.extract_text() for page in pdf.pages)
    assert "Alex Example" in text and "Half-count tag" in text and "35/6" in text
    assert abs(float(pdf.pages[0].mediabox.width) - 595.2756) < 0.1
    links = [annotation.get_object().get("/A", {}).get("/URI") for page in pdf.pages for annotation in page.get("/Annots", [])]
    assert "https://example.org/demo" in links
    assert "https://example.org/sheet" in links
    assert any(page.get("/Resources").get("/XObject") for page in pdf.pages)


def test_docx_has_readable_tables_links_qr_and_large_print_settings():
    payload = export_project(fixture_project(), "docx", large_print=True, print_qr=True)[0]
    document = Document(BytesIO(payload))
    content = "\n".join(p.text for p in document.paragraphs) + "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    assert "Evening Practice" in content and "Diagonal finish" in content and "Right step" in content
    assert document.styles["Normal"].font.size.pt == 16
    assert round(document.sections[0].page_width.inches, 2) == 8.5
    assert len(document.inline_shapes) == 2
    usable_width = document.sections[0].page_width - document.sections[0].left_margin - document.sections[0].right_margin
    for table in document.tables:
        for row in table.rows:
            assert abs(sum(cell.width for cell in row.cells) - usable_width) < 1000
            assert round(row.cells[0].width.inches, 2) == 1.4
    with zipfile.ZipFile(BytesIO(payload)) as archive:
        relationships = archive.read("word/_rels/document.xml.rels").decode()
        assert "https://example.org/demo" in relationships and "TargetMode=\"External\"" in relationships
        assert "tblHeader" in archive.read("word/document.xml").decode()
        assert "cantSplit" in archive.read("word/document.xml").decode()


def test_csv_and_xlsx_keep_exact_fractions_and_do_not_execute_formula_text():
    project = fixture_project()
    project["draft"]["sheet_meta"]["dance_title"] = '=HYPERLINK("https://example.org","formula")'
    csv_rows = list(csv.reader(StringIO(export_project(project, "csv")[0].decode("utf-8-sig"))))
    assert csv_rows[1][0].startswith("'=")
    assert csv_rows[-1][10] == "1/3"
    workbook = load_workbook(BytesIO(export_project(project, "xlsx")[0]))
    assert workbook["Choreography"]["A2"].data_type == "s"
    assert workbook["Choreography"].cell(workbook["Choreography"].max_row, 11).value == "1/3"
    assert workbook["References"]["B2"].hyperlink.target == "https://example.org/demo"


def test_subtitle_golden_times_use_same_exact_occurrence_timeline():
    project = fixture_project()
    srt = export_project(project, "srt")[0].decode()
    vtt = export_project(project, "vtt")[0].decode()
    assert "00:00:02,000 --> 00:00:02,500" in srt
    assert "00:00:04,000 --> 00:00:04,250" in srt
    assert "00:00:04,750 --> 00:00:04,917" in srt
    assert vtt.startswith("WEBVTT\n\n")
    assert "00:00:04.750 --> 00:00:04.917" in vtt
    assert srt.count(" --> ") == len(build_sheet_model(project)["compiled"]["events"])


def test_piecewise_tempo_captions_and_package_preserve_reviewed_anchor_map():
    project = fixture_project()
    anchors = [{"count": 0, "time": 2}, {"count": 2, "time": 3}, {"count": 4, "time": 5}]
    project["draft"]["music_map"]["anchors"] = anchors
    srt = export_project(project, "srt")[0].decode()
    assert "00:00:02,000 --> 00:00:02,500" in srt
    assert "00:00:03,000 --> 00:00:04,000" in srt
    assert "00:00:05,000 --> 00:00:05,500" in srt
    assert "00:00:06,500 --> 00:00:06,833" in srt
    restored = read_project_package(export_project(project, "zip")[0])
    assert restored["draft"]["music_map"]["anchors"] == anchors


@pytest.mark.parametrize("anchors", [[{"count": 2, "time": 2}, {"count": 1, "time": 3}],
                                     [{"count": 0, "time": 2}, {"count": 1, "time": 1}],
                                     [{"count": 0, "time": float("inf")}], [{"count": 0}]])
def test_caption_malformed_anchor_maps_fail_actionably(anchors):
    project = fixture_project(); project["draft"]["music_map"]["anchors"] = anchors
    with pytest.raises(ExportError, match="music timing"):
        export_project(project, "srt")


@pytest.mark.parametrize("mapping", [{}, {"bpm": 0}, {"bpm": -10}, {"bpm": "bad"}, {"bpm": 120, "first_count": -1}, {"bpm": 120, "variable_tempo": True}])
def test_caption_invalid_or_unsupported_timing_fails_actionably(mapping):
    project = fixture_project(); project["draft"]["music_map"] = mapping
    with pytest.raises(ExportError):
        export_project(project, "srt")


def test_unknown_draft_can_export_with_disclosed_status_without_verification_claim():
    project = fixture_project()
    project["draft"]["choreography"] = {"parts": [{"id": "A", "moves": [{"id": "personal", "duration_counts": "2", "text": "Personal movement"}]}]}
    text = export_project(project, "txt")[0].decode()
    assert "DRAFT - UNVERIFIED MECHANICS" in text and "Checks and review" in text


@pytest.mark.parametrize("format", ["json", "zip"])
def test_native_package_roundtrip_preserves_frozen_choreography_and_relink_metadata(format):
    project = fixture_project(); original = deepcopy(project)
    payload = export_project(project, format)[0]
    portable = read_project_package(payload)
    assert project == original
    expected = build_sheet_model(project)["compiled"]
    restored = build_sheet_model({"name": portable["name"], "song": portable["song"], "draft": portable["draft"]})["compiled"]
    assert expected["events"] == restored["events"]
    assert portable["draft"]["music_map"]["key"] == "C major"
    assert portable["media_references"][0]["filename"] == "foot-demo.mp4"
    assert portable["media_references"][0]["needs_relink"] is True
    decoded = json.dumps(portable)
    assert "DO-NOT-EXPORT" not in decoded and "PRIVATE LYRICS" not in decoded
    assert "C:\\Private" not in decoded and "unreleased.wav" not in decoded
    assert "api_key" not in decoded and "connection" not in decoded


@pytest.mark.parametrize("format", ["json", "zip"])
def test_move_source_annotations_and_media_relinks_survive_without_altering_mechanics(format):
    project=fixture_project()
    part=project["draft"]["choreography"]["parts"][0]
    part.update(notes="Part teaching paragraph.\nKeep this exact second line.", author="Synthetic author")
    move=part["moves"][0]
    move.update(snapshot_id="local-pattern@7", library_version=7, reference_fingerprint="f" * 64,
                explanation="An independently authored explanation.\nSecond paragraph.",
                notes="The rear view makes this clearer.",
                links=[{"label":"Teaching video","url":"https://example.org/teaching"},
                       {"label":"Unsafe","url":"javascript:alert(1)"}],
                origins=[{"kind":"file","label":"Authored instructions","locator":r"C:\Private\source.docx",
                          "url":"https://example.org/author","rights_note":"Synthetic fixture", "api_key":"PRIVATE-ANNOTATION-KEY"}],
                review={"status":"AUTHOR_REVIEWED","reviewer":"Synthetic author","notes":"Author review only", "api_key":"PRIVATE-ANNOTATION-KEY"},
                reference_record={"name":"Synthetic reference","description":"Exact reference paragraph.","aliases":["Test name"],
                                  "counts":1,"api_key":"PRIVATE-ANNOTATION-KEY","path":r"C:\Private\reference.json"},
                reported_mechanics={"counts":1,"net_rotation_deg":0},
                attachments=[{"id":"media-example","original_filename":"rear-demo.mp4","mime":"video/mp4",
                              "sha256":"e"*64,"caption":"Rear view","orientation":"back","start_seconds":"1/2","end_seconds":"3",
                              "path":r"C:\Private\rear-demo.mp4", "api_key":"PRIVATE-ANNOTATION-KEY"}])
    before=build_sheet_model(project)["compiled"]
    package=read_project_package(export_project(project,format)[0])
    restored_move=package["draft"]["choreography"]["parts"][0]["moves"][0]
    assert restored_move["snapshot_id"] == "local-pattern@7" and restored_move["reference_fingerprint"] == "f"*64
    assert restored_move["explanation"] == move["explanation"] and restored_move["notes"] == move["notes"]
    assert package["draft"]["choreography"]["parts"][0]["notes"] == part["notes"]
    assert restored_move["origins"][0]["locator"] == "source.docx"
    assert restored_move["reference_record"]["description"] == "Exact reference paragraph."
    media=restored_move["attachments"][0]
    assert media["id"] == "media-example" and media["sha256"] == "e"*64 and media["filename"] == "rear-demo.mp4"
    assert media["needs_relink"] is True and media["start_seconds"] == "1/2"
    assert any(item["id"] == "media-example" for item in package["media_references"])
    decoded=json.dumps(package)
    assert "PRIVATE-ANNOTATION-KEY" not in decoded and "C:\\\\Private" not in decoded and "javascript:" not in decoded
    rebuilt=build_sheet_model({"name":package["name"],"song":package["song"],"draft":package["draft"]})["compiled"]
    assert rebuilt["events"] == before["events"] and rebuilt["source_hash"] == before["source_hash"]


def test_sheet_references_include_only_safe_links_from_used_moves_and_overrides():
    project=fixture_project()
    original=project["draft"]["choreography"]["parts"][0]["moves"][0]
    original["links"]=[{"label":"Used teaching link","url":"https://example.org/used"}]
    original["origins"]=[{"label":"Author","url":"https://example.org/provenance"}]
    unused=deepcopy(project["draft"]["choreography"]["parts"][0]); unused["id"]="unused"
    unused["moves"][0]["links"]=[{"label":"Unused","url":"https://example.org/unused"}]
    project["draft"]["choreography"]["parts"].append(unused)
    override=deepcopy(original)
    override["links"]=[{"label":"Changed occurrence video","url":"https://example.org/override"},
                       {"label":"Local file","url":"file:///C:/Private/video.mp4"}]
    project["draft"]["choreography"]["routine"][0]["overrides"]={original["id"]:override}
    model=build_sheet_model(project)
    urls={link["url"] for link in model["links"]}
    assert {"https://example.org/used", "https://example.org/override", "https://example.org/provenance"} <= urls
    assert "https://example.org/unused" not in urls and not any(url.startswith("file:") for url in urls)
    html=export_project(project,"html")[0].decode()
    assert 'href="https://example.org/override"' in html
    package=read_project_package(export_project(project,"zip")[0])
    saved_override=package["draft"]["choreography"]["routine"][0]["overrides"][original["id"]]
    assert saved_override["links"] == [{"label":"Changed occurrence video","url":"https://example.org/override"}]


def test_package_manifest_hashes_exact_project_bytes():
    payload = export_project(fixture_project(), "zip")[0]
    with zipfile.ZipFile(BytesIO(payload)) as archive:
        assert set(archive.namelist()) == {"manifest.json", "project.json"}
        manifest = json.loads(archive.read("manifest.json"))
        project_bytes = archive.read("project.json")
        assert manifest["entries"] == [{"path": "project.json", "size": len(project_bytes), "sha256": hashlib.sha256(project_bytes).hexdigest()}]


def _archive(entries, compression=zipfile.ZIP_STORED):
    output = BytesIO()
    with zipfile.ZipFile(output, "w", compression=compression) as archive:
        for name, payload in entries:
            archive.writestr(name, payload)
    return output.getvalue()


@pytest.mark.parametrize("name", ["../outside.json", "/absolute.json", "C:\\outside.json", "sub/../../outside.json", "extra.txt"])
def test_archive_paths_outside_declared_manifest_are_rejected(name):
    with pytest.raises(ExportError, match="path"):
        read_project_package(_archive([(name, b"{}")] ))


def test_duplicate_members_symlinks_compression_bombs_and_hash_tampering_fail():
    with pytest.warns(UserWarning):
        duplicate = _archive([("project.json", b"{}"), ("project.json", b"{}")])
    with pytest.raises(ExportError, match="Duplicate"):
        read_project_package(duplicate)
    info = zipfile.ZipInfo("project.json"); info.create_system = 3; info.external_attr = (stat.S_IFLNK | 0o777) << 16
    with pytest.raises(ExportError, match="symlinks"):
        read_project_package(_archive([(info, b"target")]))
    bomb = _archive([("project.json", b"0" * 200000)], zipfile.ZIP_DEFLATED)
    assert len(bomb) < 200000 / MAX_COMPRESSION_RATIO
    with pytest.raises(ExportError, match="limits"):
        read_project_package(bomb)
    valid = export_project(fixture_project(), "zip")[0]
    with zipfile.ZipFile(BytesIO(valid)) as archive:
        tampered = _archive([("manifest.json", archive.read("manifest.json")), ("project.json", archive.read("project.json") + b" ")])
    with pytest.raises(ExportError, match="integrity"):
        read_project_package(tampered)


@pytest.mark.parametrize("payload", [b'{"kind":"a","kind":"b"}', b'{"bpm":NaN}', b'[]', b'not JSON'])
def test_malformed_json_does_not_restore(payload):
    with pytest.raises(ExportError):
        read_project_package(payload)


def test_future_schema_rejected_before_creating_a_project(monkeypatch, tmp_path):
    monkeypatch.setattr(store, "PROJECTS_DIR", str(tmp_path / "projects"))
    package = portable_project(fixture_project()); package["schema_version"] = 900
    with pytest.raises(ExportError, match="schema"):
        restore_project_package(json.dumps(package).encode())
    assert not (tmp_path / "projects").exists()


@pytest.mark.parametrize("location", ["package", "choreography", "manifest"])
def test_boolean_schema_versions_are_rejected(location):
    package = portable_project(fixture_project())
    if location == "manifest":
        payload = export_project(fixture_project(), "zip")[0]
        with zipfile.ZipFile(BytesIO(payload)) as archive:
            manifest = json.loads(archive.read("manifest.json")); manifest["schema_version"] = True
            payload = _archive([("manifest.json", json.dumps(manifest).encode()), ("project.json", archive.read("project.json"))])
    else:
        target = package if location == "package" else package["draft"]["choreography"]
        target["schema_version"] = True
        payload = json.dumps(package).encode()
    with pytest.raises(ExportError, match="schema"):
        read_project_package(payload)


def test_restore_creates_distinct_projects_preserving_original_and_frozen_definitions(monkeypatch, tmp_path):
    monkeypatch.setattr(store, "PROJECTS_DIR", str(tmp_path / "projects"))
    existing = store.create_project("Evening Practice", "Original song", "Original artist")
    original_bytes = (tmp_path / "projects" / existing["id"] / "project.json").read_bytes()
    payload = export_project(fixture_project(), "zip")[0]
    first, second = restore_project_package(payload), restore_project_package(payload)
    assert len({existing["id"], first["id"], second["id"]}) == 3
    assert (tmp_path / "projects" / existing["id"] / "project.json").read_bytes() == original_bytes
    saved = store.load_project(first["id"])
    assert saved["song"]["path"] is None
    assert build_sheet_model(saved)["compiled"]["total_counts"] == "35/6"
    assert saved["draft"]["attachments"][0]["needs_relink"]


def test_import_discards_injected_identity_audio_path_and_provider_fields(monkeypatch, tmp_path):
    monkeypatch.setattr(store, "PROJECTS_DIR", str(tmp_path / "projects"))
    package = portable_project(fixture_project())
    package.update(id="overwrite-me", provider_settings={"api_key": "BAD-SECRET"})
    package["song"]["path"] = r"C:\private\file.wav"
    package["draft"]["music_map"]["api_key"] = "BAD-SECRET"
    restored = restore_project_package(json.dumps(package).encode())
    saved = store.load_project(restored["id"])
    assert restored["id"] != "overwrite-me"
    assert saved["song"]["path"] is None
    assert "BAD-SECRET" not in json.dumps(saved)


def test_preview_is_read_only_and_routes_match_creator_ui(monkeypatch, tmp_path):
    monkeypatch.setattr(store, "PROJECTS_DIR", str(tmp_path / "projects"))
    payload = export_project(fixture_project(), "zip")[0]
    application = FastAPI(); application.include_router(router)
    with TestClient(application) as client:
        response = client.post("/api/creator/import/preview", files={"file": ("dance.zip", payload, "application/zip")})
        assert response.status_code == 200
        assert response.json()["name"] == "Evening Practice"
        assert not (tmp_path / "projects").exists()
        response = client.post("/api/creator/import", files={"file": ("dance.zip", payload, "application/zip")})
        assert response.status_code == 200
        pid = response.json()["id"]
        response = client.get(f"/api/creator/projects/{pid}/export/pdf?paper=a4&large_print=1")
        assert response.status_code == 200 and response.content.startswith(b"%PDF")
        assert "filename*=" in response.headers["content-disposition"]
        assert client.get(f"/api/creator/projects/{pid}/export/bogus").status_code == 400
        assert client.post("/api/creator/import", files={"file": ("bad.json", b"{}")}).status_code == 400


@pytest.mark.parametrize("value", ["javascript:alert(1)", "file:///C:/private", "https://user:secret@example.org", "https://example.org:bad", "//example.org"])
def test_unsafe_link_schemes_and_embedded_credentials_are_not_links(value):
    assert safe_url(value) is None
